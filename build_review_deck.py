# -*- coding: utf-8 -*-
"""
干部试用期转正述职 PPT 生成脚本（基于上传的讯飞干部转正述职模板）
用法：python3 build_review_deck.py <template.pptx> <out.pptx>
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else "template.pptx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "out.pptx"

# ───────────────────────── 设计令牌 ─────────────────────────
BLUE = "0C4BA1"      # 模板品牌蓝
NAVY = "0A2A5E"
BLUE2 = "3F7BD6"     # 次级蓝
TINT = "EEF3FB"      # 卡片底色
TINT2 = "DCE7F8"
TXT = "2F3E4F"
MUTE = "6F7E8F"
LINE = "D3DCE8"
GREEN = "1F8F5F"
GREEN_T = "E4F4EC"
AMBER = "C57B14"
AMBER_T = "FBF0DD"
GRAYC = "7A8898"
WHITE = "FFFFFF"

SLIDE_W, SLIDE_H = 13.333, 7.5
L, R = 0.4, 12.93
CW = R - L
CT = 1.28

prs = Presentation(TEMPLATE)
LAYOUTS = {lay.name: lay for lay in prs.slide_layouts}
LAY_CONTENT = LAYOUTS["4_标题和内容"]


def rgb(h):
    return RGBColor.from_string(h)


# ───────────────────────── 基础绘制 ─────────────────────────
def _parse_runs(s):
    """'**粗体**普通' → [(text, bold)]"""
    parts = s.split("**")
    runs = []
    for i, part in enumerate(parts):
        if part:
            runs.append((part, i % 2 == 1))
    return runs


def _norm_paras(content):
    if isinstance(content, str):
        return [{"text": t} for t in content.split("\n")]
    if isinstance(content, tuple) and len(content) == 2 and isinstance(content[1], dict):
        content = [content]
    out = []
    for item in content:
        if isinstance(item, str):
            out.append({"text": item})
        elif isinstance(item, tuple):
            d = dict(item[1])
            d["text"] = item[0]
            out.append(d)
        else:
            out.append(dict(item))
    return out


def _bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.16))))
    pPr.set("indent", str(-int(Inches(0.16))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    c = etree.SubElement(buClr, qn("a:srgbClr"))
    c.set("val", color)
    buSz = etree.SubElement(pPr, qn("a:buSzPct"))
    buSz.set("val", "100000")
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", "•")


def fill_tf(tf, content, size=11, color=TXT, bold=False, align="l", valign="t",
            margin=(0.06, 0.04, 0.06, 0.04), ls=1.12, sa=3, bullet=False, bcolor=None):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[valign]
    paras = _norm_paras(content)
    # 清空已有段落
    txBody = tf._txBody
    for p in list(txBody.findall(qn("a:p"))):
        txBody.remove(p)
    for pd in paras:
        p = tf.add_paragraph()
        al = pd.get("align", align)
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
        p.line_spacing = pd.get("ls", ls)
        p.space_after = Pt(pd.get("sa", sa))
        if "sb" in pd:
            p.space_before = Pt(pd["sb"])
        psize = pd.get("size", size)
        pcolor = pd.get("color", color)
        pbold = pd.get("bold", bold)
        for text, b in _parse_runs(pd["text"]):
            r = p.add_run()
            r.text = text
            r.font.size = Pt(psize)
            r.font.bold = bool(pbold or b)
            r.font.color.rgb = rgb(pcolor)
        if pd.get("bullet", bullet):
            _bullet(p, pd.get("bcolor", bcolor or BLUE))
    return tf


def text(slide, x, y, w, h, content, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, content, **kw)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, radius=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def card(slide, x, y, w, h, fill=TINT, line=None, radius=0.05):
    return box(slide, x, y, w, h, fill=fill, line=line, radius=radius, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def chip(slide, x, y, w, h, label, fill=BLUE, color=WHITE, size=10, bold=True, line=None, radius=0.5):
    shp = box(slide, x, y, w, h, fill=fill, line=line, radius=radius, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    fill_tf(shp.text_frame, label, size=size, color=color, bold=bold, align="c", valign="m",
            margin=(0.04, 0.0, 0.04, 0.0), sa=0)
    return shp


def dot(slide, x, y, d, label, fill=BLUE, color=WHITE, size=11):
    shp = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    fill_tf(shp.text_frame, label, size=size, color=color, bold=True, align="c", valign="m",
            margin=(0, 0, 0, 0), sa=0)
    return shp


def chevron(slide, x, y, w=0.16, h=0.32, fill=BLUE2):
    return box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.CHEVRON)


def bar(slide, y, content, h=0.55, fill=NAVY, size=11.5, label=None):
    """页底结论条"""
    card(slide, L, y, CW, h, fill=fill, radius=0.12)
    if label:
        chip(slide, L + 0.15, y + (h - 0.3) / 2, 1.0, 0.3, label, fill=WHITE, color=fill, size=9.5)
        text(slide, L + 1.25, y, CW - 1.4, h, content, size=size, color=WHITE, bold=True, valign="m")
    else:
        text(slide, L + 0.2, y, CW - 0.4, h, content, size=size, color=WHITE, bold=True, valign="m")


def note(slide, y, content, h=0.36):
    text(slide, L, y, CW, h, content, size=9.5, color=MUTE, valign="m")


def stat(slide, x, y, w, h, big, label, fill=TINT, big_color=BLUE):
    card(slide, x, y, w, h, fill=fill)
    text(slide, x, y + 0.06, w, h * 0.55, big, size=20, color=big_color, bold=True, align="c", valign="m",
         margin=(0.04, 0, 0.04, 0), sa=0)
    text(slide, x, y + h * 0.56, w, h * 0.42, label, size=9.5, color=MUTE, align="c", valign="t",
         margin=(0.06, 0, 0.06, 0), sa=0)


def layer_card(slide, x, y, w, h, label, color, bullets, size=10.5, label_fill=None):
    card(slide, x, y, w, h)
    chip(slide, x + 0.14, y + 0.14, 1.25, 0.32, label, fill=label_fill or color, size=10)
    text(slide, x + 0.1, y + 0.56, w - 0.2, h - 0.62, [(b, {"bullet": True, "bcolor": color}) for b in bullets],
         size=size, sa=4)


def case_layers(slide, y, h, layers, size=10.5):
    """四层案例：障碍→判断→行动→变化"""
    n = len(layers)
    gap = 0.24
    w = (CW - gap * (n - 1)) / n
    for i, (label, color, bullets) in enumerate(layers):
        x = L + i * (w + gap)
        layer_card(slide, x, y, w, h, label, color, bullets, size=size)
        if i < n - 1:
            chevron(slide, x + w + 0.04, y + 0.16, fill=color)


def titled_card(slide, x, y, w, h, title, bullets, num=None, size=10.5, title_size=12.5, fill=TINT):
    card(slide, x, y, w, h, fill=fill)
    tx = x + 0.16
    if num is not None:
        dot(slide, x + 0.16, y + 0.16, 0.34, str(num), size=11)
        tx = x + 0.6
    text(slide, tx, y + 0.12, w - (tx - x) - 0.1, 0.42, title, size=title_size, color=NAVY, bold=True, valign="m",
         margin=(0.02, 0, 0.02, 0), sa=0)
    text(slide, x + 0.12, y + 0.6, w - 0.24, h - 0.66, [(b, {"bullet": True}) for b in bullets], size=size, sa=4)


def band(slide, x, y, w, h, label, bullets, cols=2, label_w=1.5, color=BLUE, size=10.5):
    """横向条带：左侧标签，右侧分栏要点"""
    card(slide, x, y, w, h)
    chip(slide, x + 0.14, y + (h - 0.36) / 2, label_w, 0.36, label, fill=color, size=10)
    cx = x + label_w + 0.35
    cw = (w - label_w - 0.5)
    if cols == 1:
        text(slide, cx, y + 0.08, cw, h - 0.14, [(b, {"bullet": True}) for b in bullets], size=size, sa=4, valign="m")
    else:
        half = (len(bullets) + 1) // 2
        colw = (cw - 0.2) / 2
        text(slide, cx, y + 0.08, colw, h - 0.14, [(b, {"bullet": True}) for b in bullets[:half]], size=size, sa=4, valign="m")
        text(slide, cx + colw + 0.2, y + 0.08, colw, h - 0.14, [(b, {"bullet": True}) for b in bullets[half:]], size=size, sa=4, valign="m")


def cell_borders(cell, bottom=LINE, top=None, left=None, right=None, w=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    specs = [("a:lnL", left), ("a:lnR", right), ("a:lnT", top), ("a:lnB", bottom)]
    for i, (tag, color) in enumerate(specs):
        ln = etree.Element(qn(tag))
        if color:
            ln.set("w", str(int(Pt(w))))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            c = etree.SubElement(sf, qn("a:srgbClr"))
            c.set("val", color)
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "solid")
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)


def table(slide, x, y, colw, header, rows, row_h=0.8, header_h=0.42, size=10, header_size=10.5,
          first_col_bold=True, zebra=True, valign="m", aligns=None, header_fill=BLUE, cell_pad=0.07):
    nrows = len(rows) + 1
    ncols = len(colw)
    total_w = sum(colw)
    total_h = header_h + row_h * len(rows)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(total_w), Inches(total_h))
    tbl = gf.table
    # 去掉默认表格样式
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(colw):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, nrows):
        tbl.rows[r].height = Inches(row_h)
    for c, h in enumerate(header):
        cell = tbl.cell(0, c)
        cell_borders(cell, bottom=None)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(header_fill)
        cell.margin_left = cell.margin_right = Inches(cell_pad)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_tf(cell.text_frame, h, size=header_size, color=WHITE, bold=True, align="l", valign="m",
                margin=(cell_pad, 0.02, cell_pad, 0.02), sa=0)
    for r, row in enumerate(rows, start=1):
        rf = (WHITE if (r % 2 == 1 or not zebra) else "F6F8FC")
        for c, content in enumerate(row):
            cell = tbl.cell(r, c)
            cell_borders(cell, bottom=LINE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(rf)
            cell.margin_left = cell.margin_right = Inches(cell_pad)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[valign]
            al = (aligns[c] if aligns else "l")
            bold = first_col_bold and c == 0
            color = NAVY if (first_col_bold and c == 0) else TXT
            fill_tf(cell.text_frame, content, size=size, color=color, bold=bold, align=al, valign=valign,
                    margin=(cell_pad, 0.03, cell_pad, 0.03), sa=2, ls=1.08)
    return tbl


def role_tag(slide, label):
    w = min(3.2, 0.32 + 0.16 * len(label) + 0.3)
    chip(slide, R - w, 0.6, w, 0.36, label, fill=WHITE, color=BLUE, size=10, line=BLUE)


def page_no(slide, n):
    text(slide, 12.0, 7.02, 0.95, 0.3, str(n), size=9, color=MUTE, align="r", valign="m", margin=(0, 0, 0.05, 0), sa=0)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


PAGE = [2]  # 当前页码计数（P1、P2 为模板原页）


def content_slide(tag, title, role=None):
    s = prs.slides.add_slide(LAY_CONTENT)
    PAGE[0] += 1
    to_remove = []
    for ph in s.placeholders:
        y = ph.top / 914400
        if ph.placeholder_format.type == 1 or "标题" in ph.name:  # TITLE
            fill_tf(ph.text_frame, title, size=20, color=BLUE, bold=True, valign="m", margin=(0.05, 0, 0.05, 0), sa=0)
        elif y < 0.5:
            fill_tf(ph.text_frame, tag, size=14, color=WHITE, bold=True, valign="m", margin=(0.1, 0, 0.05, 0), sa=0)
        else:
            to_remove.append(ph)
    for ph in to_remove:
        ph._element.getparent().remove(ph._element)
    if role:
        role_tag(s, role)
    page_no(s, PAGE[0])
    return s


# ───────────────────────── 模板原页处理 ─────────────────────────
def replace_text(shape, lines, size=None, color=None, bold=None):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    rpr = copy.deepcopy(p0.runs[0]._r.get_or_add_rPr()) if p0.runs else None
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    txBody = tf._txBody
    for p in list(txBody.findall(qn("a:p"))):
        txBody.remove(p)
    for line in lines:
        p = etree.SubElement(txBody, qn("a:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn("a:r"))
        if rpr is not None:
            rp = copy.deepcopy(rpr)
            if size:
                rp.set("sz", str(int(size * 100)))
            r.append(rp)
        t = etree.SubElement(r, qn("a:t"))
        t.text = line


slides0 = list(prs.slides)
cover, toc, end = slides0[0], slides0[1], slides0[9]

# 封面
tbs = [sh for sh in cover.shapes if sh.has_text_frame]
title_box = min(tbs, key=lambda s: s.top)
who_box = max(tbs, key=lambda s: s.top)
replace_text(title_box, ["干部试用期转正述职"])
text(cover, 1.94, 3.28, 9.45, 0.6, "国家级战略业务突破与国际咨询能力建设", size=22, color=NAVY, bold=False, align="c", valign="m", sa=0)
who_box.left, who_box.width = Inches(3.67), Inches(6.0)
who_box.top, who_box.height = Inches(4.95), Inches(1.2)
replace_text(who_box, ["汇报人：陈昊　市场咨询总监", "2026年9月"])
notes(cover, "开场一句：这段时间，我一方面在巴西亲自牵引国家AI项目和重点应用合作，另一方面同步建设国际咨询队伍、持续支撑各区域业务。本次汇报两条线取得的成果、关键贡献，以及下一阶段如何完成商业闭环和强化组织承接。")

# 目录
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh

label_map = {"管理角色澄清": "管理角色澄清与成果回顾", "Part5": "Part 5", "Part6": "Part 6"}
for sh in walk(toc.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() in label_map:
        replace_text(sh, [label_map[sh.text_frame.text.strip()]])
text(toc, 6.2, 1.84, 6.3, 0.42,
     "战略目标与核心成果 · 管理角色澄清表 · 双线并发 · 成果一：巴西战略业务突破 · 成果二：国际咨询能力与全球支撑 · 角色自评",
     size=10, color=MUTE, valign="t", sa=0)

# 结尾
text(end, 2.2, 4.45, 8.93, 1.7, [
    ("业务上　持续牵引巴西国家AI与重点应用合作，取得重大阶段突破并推进商业落地", {"align": "c"}),
    ("组织上　建设国际咨询队伍，形成多区域方案、商机与风险管理的实际支撑", {"align": "c"}),
    ("下一阶段　将阶段成果转化为可持续商业结果，将关键个人经验转化为稳定组织能力", {"align": "c"}),
], size=13, color=NAVY, valign="t", sa=8)
notes(end, "结尾只保留两条成果和一个后续方向，不再重复项目清单。")

# ───────────────────────── P3 战略目标承接 ─────────────────────────
s = content_slide("1-管理角色澄清", "战略目标承接：两条成果线的结果与本人作用", role="看定位")
# 顶部两问
card(s, L, CT, CW, 0.5, fill=TINT2, radius=0.3)
chip(s, L + 0.12, CT + 0.09, 1.15, 0.32, "可衡量结果", size=9.5)
text(s, L + 1.32, CT, 4.95, 0.5, "两条成果线：巴西战略业务突破；国际咨询队伍与多区域支撑", size=10.5, color=NAVY, bold=True, valign="m", sa=0)
chip(s, L + 6.45, CT + 0.09, 1.0, 0.32, "实现策略", size=9.5)
text(s, L + 7.55, CT, 4.9, 0.5, "重大项目亲自攻坚，咨询组织同步建设，两线并行推进", size=10.5, color=NAVY, bold=True, valign="m", sa=0)

def outcome_card(x, y, w, h, head, blocks):
    card(s, x, y, w, h, fill=WHITE, line=LINE)
    box(s, x, y, w, 0.52, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    box(s, x, y + 0.3, w, 0.22, fill=BLUE)  # 拉直下沿
    text(s, x + 0.15, y, w - 0.3, 0.52, head, size=12.5, color=WHITE, bold=True, valign="m", sa=0)
    bh = (h - 0.52 - 0.2) / 3
    for i, (label, color, bullets) in enumerate(blocks):
        by = y + 0.62 + i * bh
        chip(s, x + 0.15, by + 0.06, 1.0, 0.3, label, fill=color, size=9.5)
        text(s, x + 1.28, by, w - 1.4, bh, [(b, {"bullet": True, "bcolor": color}) for b in bullets], size=10.5, sa=3)
        if i < 2:
            box(s, x + 0.15, by + bh - 0.03, w - 0.3, 0.01, fill=LINE)

outcome_card(L, 1.92, 6.12, 4.12, "成果一　牵引巴西战略业务取得重大阶段突破，持续推进商业落地", [
    ("公司所得", BLUE, ["国家AI合作位置明确：总统公开宣布相关合作、三方MoU签署",
                    "主要技术路线获客户认可，核心内容进入采购需求（TR）",
                    "采购合同与本地承接条件持续推进；SERPRO招聘应用进入PoC开发与部署"]),
    ("本人作用", NAVY, ["承担讯飞侧总体推进，补位客户与销售工作",
                    "亲自完成核心英语咨询与关键技术澄清",
                    "组织政府、华为、商法财、交付多方协同"]),
    ("实现条件", GRAYC, ["客户需求与交易模式持续变化",
                     "前期销售承接与本地组织尚不完善",
                     "关键外部主体不受本人直接管理"]),
])
outcome_card(L + 6.41, 1.92, 6.12, 4.12, "成果二　从0到1建设国际咨询队伍，形成跨区域支撑能力", [
    ("公司所得", BLUE, ["重点区域形成主责咨询配置（Q1记录：6人、5个区域）",
                    "持续产出国别方案、HLD、SoC、RFI、可研、报价与合同范围",
                    "骨干已能承担重要项目主责"]),
    ("本人作用", NAVY, ["完成选人、培养、部署与关键质量把关",
                    "参与跨区域竞争策略、范围、利润与交付风险判断"]),
    ("实现条件", GRAYC, ["本人同期深度投入巴西",
                     "团队仍在成长；多国并发需现场支撑与轮换"]),
])
bar(s, 6.2, "不是在成熟组织里完成单一专业任务，而是在组织建设与业务推进并行的阶段，既亲自承担关键突破，也持续组织团队形成多区域产出。", h=0.6, label="总判断", size=11.5)
notes(s, "主讲2分钟。先讲公司拿到了什么，再讲本人起了什么作用，最后讲是在什么条件下完成的。全篇只在这一页评价自己，后面全部用事实。")

# ───────────────────────── P4 管理角色澄清表 ─────────────────────────
s = content_slide("1-管理角色澄清", "管理角色澄清表：四类角色的目标、投入与自评")
table(s, L, CT,
      [1.6, 2.55, 0.6, 2.7, 2.35, 0.75, 1.98],
      ["管理角色", "目标（业绩结果）", "权重", "高回报活动（3项）", "管理机制", "时间投入", "完成情况自评"],
      [
          [["**岗位特性类**", ("关键客户咨询与复杂方案牵引", {"size": 9, "color": MUTE})],
           ["① 巴西主要技术路线获客户认可并进入采购需求（TR）", "② 重点区域形成国别方案与技术建议"],
           "—",
           ["① 亲自承担RNP关键英语咨询与技术澄清", "② 国别方案与技术路线判断", "③ 重大效果、范围、权利边界把关"],
           ["关键材料评审", "问题书面化", "面向团队的标准口径材料"],
           "—",
           [("达成", {"bold": True, "color": GREEN}), "案例见 P8、P9、P20"]],
          [["**目标管理类**", ("巴西重点项目总体推进与全球咨询目标统筹", {"size": 9, "color": MUTE})],
           ["① RNP进入采购流程并具备签约条件", "② SERPRO进入PoC验证", "③ 各区域商机按成熟度分级推进"],
           "—",
           ["① 里程碑与卡点识别", "② 政府、伙伴与内部决策输入", "③ 商机取舍与优先级"],
           ["项目周报与向上汇报", "现场LTC", "决策输入回传一线"],
           "—",
           [("阶段达成", {"bold": True, "color": AMBER}), "商业闭环未完成", "见 P7、P10、P13、P21"]],
          [["**人与团队类**", ("招募、培养、部署咨询骨干并支持销售承接", {"size": 9, "color": MUTE})],
           ["① 重点区域形成主责咨询配置", "② 骨干能独立承接项目主责"],
           "—",
           ["① 实战带教", "② 集中办公与现场轮换", "③ 主责交接与销售导入"],
           ["集中办公", "实战复盘", "主备安排"],
           "—",
           [("形成中", {"bold": True, "color": AMBER}), "见 P15、P16"]],
          [["**运营类**", ("跨部门、跨伙伴协同与交付风险前置", {"size": 9, "color": MUTE})],
           ["① 售前、合同、交付之间不出现断层", "② 本地承接条件明确"],
           "—",
           ["① SoC与合同范围评审", "② 法财与交付前置", "③ 协同机制建设"],
           ["采购会议", "书面求证", "三层沟通通道；材料复核"],
           "—",
           [("达成", {"bold": True, "color": GREEN}), "标准化不足", "见 P11、P12、P17、P24"]],
      ], row_h=1.12, size=9.5, header_size=10, valign="t", aligns=["l", "l", "c", "l", "l", "c", "l"], first_col_bold=False)
note(s, 6.35, "权重与时间投入比例由本人确认后填写；完成情况自评的依据见对应页码。")
notes(s, "主讲1分钟。只讲四行角色定位，不念表格。评委关心的看定位、看结果、看过程、看时间精力分别在P3、P6、案例页和P5。")

# ───────────────────────── P5 双线并发 ─────────────────────────
s = content_slide("1-管理角色澄清", "双线并发：深度承担巴西的同时，持续组织全球咨询支撑", role="看时间精力")
colw = [1.15, 5.58, 5.58]
xs = [L, L + colw[0] + 0.11, L + colw[0] + colw[1] + 0.22]
heads = ["时段", "巴西线（本人亲自牵引）", "全球咨询线（本人组织与把关）"]
for i in range(3):
    box(s, xs[i], CT, colw[i], 0.4, fill=BLUE if i else NAVY)
    text(s, xs[i], CT, colw[i], 0.4, heads[i], size=10.5, color=WHITE, bold=True, valign="m", align="c" if i == 0 else "l")
rows = [
    ("4—5月", ["RNP方案与4批技术澄清", "SERPRO医疗、教育、人事管理场景推进", "新到现场销售的项目与客户关系导入"],
     ["埃及客户交流与售前向交付交底", "加纳、IGNITE、SDEC方案工作", "处理咨询离职、双现场驻场与轮换"]),
    ("6月", ["总集、服务化、分签等交易模式讨论", "报价刷新与付款里程碑设计", "SGP招聘需求确认与PoC启动"],
     ["埃及合同初版评审", "非洲多国RFI、概算与可研", "SDEC现场支撑；新咨询人员带教"]),
    ("8月", ["政府间文件、采购接口、联合SOW", "本地实体、律师资源与LTC现场", "SERPRO部署与验收条件沟通"],
     ["KNG范围与交付风险评审", "SDEC五个子项目SoC", "哈萨克斯坦、泰国方案与客户交流"]),
]
rh = 1.08
for r, (period, a, b) in enumerate(rows):
    y = CT + 0.48 + r * (rh + 0.08)
    card(s, xs[0], y, colw[0], rh, fill=TINT2)
    text(s, xs[0], y, colw[0], rh, period, size=13, color=NAVY, bold=True, align="c", valign="m", sa=0)
    for i, items in ((1, a), (2, b)):
        card(s, xs[i], y, colw[i], rh, fill=TINT)
        text(s, xs[i] + 0.1, y + 0.05, colw[i] - 0.2, rh - 0.1, [(t, {"bullet": True}) for t in items], size=10.5, valign="m", sa=2)
sy = CT + 0.48 + 3 * (rh + 0.08) + 0.08
sw = (CW - 3 * 0.2) / 4
for i, (big, lab) in enumerate([("5次", "七个月内中巴往返"), ("11小时", "时差下的两地协同"), ("2地", "巴西、埃及同时需常驻与轮换"), ("本人接续", "现场人员变化时跨区支援")]):
    stat(s, L + i * (sw + 0.2), sy, sw, 0.82, big, lab)
text(s, L, sy + 0.9, CW, 0.4, "难度不只是巴西项目复杂：在深度承担巴西的同时，需要同步作跨区域人员、方案与优先级安排。", size=11, color=NAVY, bold=True, valign="m", sa=0)
notes(s, "主讲1分钟。强调的是同一时期两条线并行，不是先后完成。差旅和时差只作为条件说明。")

# ───────────────────────── P6 巴西全景 ─────────────────────────
s = content_slide("1-管理角色澄清", "巴西成果全景：从合作意向到国家级项目、应用验证与本地承接", role="成果一 ｜ 看结果")
c0, c1, c2 = 2.0, 4.55, 5.5
x0, x1, x2 = L, L + c0 + 0.12, L + c0 + c1 + 0.24 + 0.24
for x, w, h_, f in ((x0, c0, "推进方向", NAVY), (x1, c1, "初始状态", GRAYC), (x2, c2, "已取得的阶段变化", BLUE)):
    box(s, x, CT, w, 0.4, fill=f)
    text(s, x, CT, w, 0.4, h_, size=10.5, color=WHITE, bold=True, valign="m", align="l")
rows = [
    ("国家AI项目", "联合路线需争取认可；实施主体、技术范围与交易模式待明确", "RNP承担项目推进；主要技术路线获认可；总统公开宣布相关合作；三方MoU签署"),
    ("技术与采购准备", "高层意向尚需转成可执行方案", "HLD、多批澄清、TR分析、SoC、报价与付款里程碑逐步形成"),
    ("政府合作支持", "政府间文件及签署授权等前置问题待解", "专题汇报、紧急需求报告、政府间文本往返进入实质推进"),
    ("SERPRO应用合作", "高层合作意愿与多场景探索起步", "医疗、教育方案交流；MGI-SGP招聘场景获选PoC并形成开发成果"),
    ("本地承接", "缺少成熟本地实体与专业支撑", "注册与经营范围准备、律师资源、LTC团队及采购接口逐步建立"),
]
rh = 0.8
for r, (a, b, c) in enumerate(rows):
    y = CT + 0.48 + r * (rh + 0.08)
    card(s, x0, y, c0, rh, fill=TINT2)
    text(s, x0 + 0.08, y, c0 - 0.16, rh, a, size=11, color=NAVY, bold=True, valign="m", sa=0)
    card(s, x1, y, c1, rh, fill="F4F6F9")
    text(s, x1 + 0.1, y, c1 - 0.2, rh, b, size=10.5, color=MUTE, valign="m", sa=0)
    chevron(s, x1 + c1 + 0.04, y + (rh - 0.32) / 2, fill=BLUE2)
    card(s, x2, y, c2, rh, fill=TINT)
    text(s, x2 + 0.1, y, c2 - 0.2, rh, c, size=10.5, color=TXT, valign="m", sa=0)
note(s, 6.22, "状态边界：正式商业合同、回款、最终验收为未完成事项，分别跟踪，不与上述阶段成果混同。数据按述职当天状态刷新。", h=0.5)
notes(s, "主讲1分钟。讲状态变化，不讲每次会议。最后一句明确边界：合同、回款、验收未完成。")

# ───────────────────────── P7 案例① 政府 ─────────────────────────
s = content_slide("1-管理角色澄清", "关键案例①　政府合作攻坚：把签约障碍转化为可决策的具体事项", role="目标管理类 ｜ 跨层级推动")
case_layers(s, CT, 2.3, [
    ("原来的障碍", GRAYC, ["巴方提出直接采购需要政府间文件支持", "使馆签署又涉及国内部委会商与授权", "单靠供应商合同谈判无法解决"]),
    ("关键判断", BLUE2, ["把前线障碍整理成高层可处理的决策事项，而非一般性介绍项目", "伙伴内部协调受限时，讯飞侧单独提交、继续推进"]),
    ("本人行动", BLUE, ["推动8月19日再次向大使专题汇报，聚焦政府间协议、合规程序与签约时间", "当晚组织现场团队完成《当前进展和紧急需求的汇报》提交发改委", "后续持续提供项目具体信息"]),
    ("可见变化", GREEN, ["8月20日反馈原则同意签署", "8月21日文本发送巴西科创部", "8月25日巴方反馈意见，进入文本确认"]),
], size=11)
# 时间轴
ty = CT + 2.55
nodes = [("8.19", "大使专题汇报"), ("当晚", "紧急需求报告提交发改委"), ("8.20", "原则同意签署"), ("8.21", "文本发送科创部"), ("8.25", "巴方反馈意见")]
box(s, L + 0.6, ty + 0.5, CW - 1.2, 0.03, fill=TINT2)
step = (CW - 1.2) / (len(nodes) - 1)
for i, (d, lab) in enumerate(nodes):
    cx = L + 0.6 + i * step
    dot(s, cx - 0.14, ty + 0.375, 0.28, "", fill=BLUE if i < 2 else GREEN)
    text(s, cx - 1.1, ty + 0.02, 2.2, 0.32, d, size=11, color=NAVY, bold=True, align="c", valign="m", sa=0)
    text(s, cx - 1.1, ty + 0.72, 2.2, 0.36, lab, size=9.5, color=TXT, align="c", valign="t", sa=0)
card(s, L, ty + 1.25, CW, 0.6, fill=TINT)
chip(s, L + 0.14, ty + 1.39, 1.0, 0.32, "同期对接", fill=NAVY, size=9.5)
text(s, L + 1.25, ty + 1.25, CW - 1.4, 0.6, "7月大使沟通 · MGI高层交流 · 8月总统府、外交部、科创部对接，分别对应文件、节奏、注册与本地经营承诺等实际议题", size=10.5, valign="m", sa=0)
note(s, ty + 1.95, "口径：后续政府决定按各方归属表述；原则同意不等于协议已签署。")
notes(s, "主讲2分钟。四层：障碍是什么、我怎么判断、我做了什么、发生了什么变化。不把政府批准归因个人，但写足组织材料、提出诉求、持续跟踪的责任。")

# ───────────────────────── P8 案例② 技术路线 ─────────────────────────
s = content_slide("1-管理角色澄清", "关键案例②　技术路线：从“自主从零训练”诉求到可交付的能力建设路径", role="岗位特性类 ｜ 专业判断")
case_layers(s, CT, 2.7, [
    ("原来的障碍", GRAYC, ["客户希望掌握自主训练能力，提出从小模型开始、从零训练", "项目转向RNP为实施主体，方案需与华为融合并首次正式面客"]),
    ("关键判断", BLUE2, ["科研探索与需要稳定效果的公共服务交付，不是同一目标", "当前交付须同时考虑效果责任、投入与可执行性", "先在成熟底座上持续训练与知识转移，再逐步探索"]),
    ("本人行动", BLUE, ["亲自承担关键英语技术与业务沟通", "联合团队围绕HLD应答4批30多个问题，持续更新书面材料", "对新增DLM、开放基础权重要求继续评审，不无条件接受"]),
    ("可见变化", GREEN, ["5月5日会议，客户明确复述“可持续训练演进，从零训练是另一层难度”", "5月周报：客户总体接受方案与技术路线，本阶段技术澄清关闭", "5月15日TR草稿与HLD及澄清回复基本一致"]),
], size=11)
sy = CT + 2.92
sw = (CW - 3 * 0.2) / 4
for i, (big, lab) in enumerate([("4批 / 30+", "客户问题联合书面应答"), ("5.5", "客户复述的理解变化"), ("5月", "总体接受方案与技术路线"), ("5.15", "TR草稿与HLD基本一致")]):
    stat(s, L + i * (sw + 0.2), sy, sw, 0.82, big, lab)
card(s, L, sy + 0.98, CW, 0.62, fill=TINT)
chip(s, L + 0.14, sy + 1.13, 0.8, 0.32, "归因", fill=NAVY, size=9.5)
text(s, L + 1.05, sy + 0.98, CW - 1.2, 0.62, "阶段认可是多轮联合工作的共同结果，不归于单次发言；技术说明不写成未经测试的性能保证，新增内容继续管理。", size=10.5, valign="m", sa=0)
notes(s, "主讲2分钟。这页证明的是专业判断：不只懂技术，而是把客户理想目标、供应商责任和实施路径组织成可讨论的选择。客户复述是直接证据，再接周报和TR。")

# ───────────────────────── P9 问答表 ─────────────────────────
s = content_slide("1-管理角色澄清", "关键咨询问答：五个核心问题的澄清及其对项目的意义", role="岗位特性类 ｜ 备讲")
table(s, L, CT, [2.85, 5.75, 3.93], ["客户核心问题", "亲自完成的关键澄清", "对项目的意义"], [
    ["巴西语料从哪里来，缺数据怎么办？", "先调查本地机构已有资源，评估数量、类型、质量，再选择购买、清洗增强或新建；语料团队与训练团队并行", "把“客户负责数据”细化为可启动的调查与建设安排"],
    ["为何不是直接从零训练？", "区分科研探索与需要稳定效果的公共服务交付；先基于成熟底座增训与知识转移，再逐步探索", "回应自主能力目标，同时明确当前项目的结果责任"],
    ["模型能否跨昇腾、英伟达生态运行？", "说明迁移、算子与软件栈适配、性能优化工作；区分基础支持与额外服务", "不用一句“兼容”掩盖工程投入与商业边界"],
    ["怎样证明葡语与本土能力改善？", "区分语言能力、本土知识与价值相关内容；提出评测基准初稿、客户讨论与量化验收口径", "把主观效果诉求转成可共同确认的评测安排"],
    ["自主使用与核心知识产权怎样兼顾？", "区分使用、定制、再分发与核心代码等权利；说明持续训练与受控数据使用安排", "把抽象“主权”要求拆成具体权利与责任讨论"],
], row_h=0.86, size=10, first_col_bold=True)
note(s, 6.15, "5月5日会议中有对应发言；5月8日语料规模、数据与源码问题为补充案例。口头讨论不写成合同条款，技术说明不写成性能保证。", h=0.5)
notes(s, "备讲页，追问技术细节时使用。")

# ───────────────────────── P10 商业模式 ─────────────────────────
s = content_slide("1-管理角色澄清", "商业模式与价值争取：从“方案可讲”推进到“交易如何成立”", role="目标管理类")
cw2 = (CW - 0.24) / 2
ch = 1.95
items = [
    ("交易模式持续调整", ["围绕联合体、讯飞总集、服务化采购、分签等路径", "讨论签约主体、资产持有、资金安排、付款与风险分担"]),
    ("客户诉求回传决策，原则回传一线", ["6月周报：OPEX模式下讯飞担任总集，完成多轮内部讨论与汇报", "高层确认的态度与原则反馈华为一线"]),
    ("明确讯飞在联合方案中的位置", ["参与华为云MA集成方案讨论", "阶段性确认训练工具链采用讯飞；推理、容器等界面继续核定"]),
    ("组织具体商业材料", ["分期实施与一期份额、回款目标讨论", "按范围刷新报价，形成付款里程碑并提交客户", "6月22—28日周报记有“48M DDP”报价与付款设计（币种未明前不换算）"]),
]
for i, (t, b) in enumerate(items):
    x = L + (i % 2) * (cw2 + 0.24)
    y = CT + (i // 2) * (ch + 0.2)
    titled_card(s, x, y, cw2, ch, t, b, num=i + 1, size=11)
bar(s, CT + 2 * ch + 0.45, "把技术范围、联合分工、公司价值与回款节奏放在一起考虑，而不只围绕客户是否喜欢方案。", h=0.55)
notes(s, "备讲或主讲半分钟。要点：我处理的是交易怎样成立，不是方案讲得好不好。")

# ───────────────────────── P11 采购与合同 ─────────────────────────
s = content_slide("1-管理角色澄清", "采购与合同攻坚：把未知条件逐项明确，形成书面推进依据", role="运营类")
c0, c1, c2 = 2.3, 7.35, 2.64
x0, x1, x2 = L, L + c0 + 0.12, L + c0 + c1 + 0.24
for x, w, h_, f in ((x0, c0, "问题", NAVY), (x1, c1, "工作", BLUE), (x2, c2, "结果状态", BLUE)):
    box(s, x, CT, w, 0.4, fill=f)
    text(s, x, CT, w, 0.4, h_, size=10.5, color=WHITE, bold=True, valign="m")
rows = [
    ("缺少本地合同专业支撑", "现场调研交流三家律所，选定与RNP沟通所需的葡语合同顾问律师", "已完成", GREEN, GREEN_T),
    ("采购前提与签约条件不明", "8月4日首次与采购部门直接对接，拉通财法与律所参与，围绕本地实体资格、签约前提与下一步时间表讨论", "前提清单形成", GREEN, GREEN_T),
    ("准备工作能否并行", "纪要明确签约准备可先行，正式签约需激活CNPJ；草稿提案作为RNP理解范围、费用归类与编制TR的输入", "并行推进中", AMBER, AMBER_T),
    ("风险事项停留在口头", "会后正式邮件对分包报备、履约担保、用工逐项求证；建立项目负责人、采购组织、执行接口三层通道；追踪11项遗留问题", "持续跟踪", AMBER, AMBER_T),
    ("伙伴正式方案未提交", "两次与华为项目负责人深入沟通，推动联合SOW攻关，同步完成内部任务分解", "推进中", AMBER, AMBER_T),
]
rh = 0.8
for r, (a, b, st, sc, sf) in enumerate(rows):
    y = CT + 0.48 + r * (rh + 0.08)
    card(s, x0, y, c0, rh, fill=TINT2)
    text(s, x0 + 0.08, y, c0 - 0.16, rh, a, size=10.5, color=NAVY, bold=True, valign="m", sa=0)
    card(s, x1, y, c1, rh, fill=TINT)
    text(s, x1 + 0.1, y, c1 - 0.2, rh, b, size=10.5, valign="m", sa=0)
    card(s, x2, y, c2, rh, fill=sf)
    chip(s, x2 + 0.3, y + (rh - 0.34) / 2, c2 - 0.6, 0.34, st, fill=sc, size=10)
note(s, 6.25, "把模糊的“等待采购”转为具体前提、材料、接口、责任与待决问题；采购会议按8月4日计一次。", h=0.4)
notes(s, "备讲页。追问采购进展时用。")

# ───────────────────────── P12 案例③ 本地承接 ─────────────────────────
s = content_slide("1-管理角色澄清", "关键案例③　本地承接：主体经营范围与方案结构的矛盾", role="运营类 ｜ 跨专业落地")
case_layers(s, CT, 3.25, [
    ("原来的障碍", GRAYC, ["采购讨论中RNP提出服务代码问题", "纯服务注册暂不含贸易与设备租赁，现有安全设备OPEX安排不能直接纳入当前签约结构", "CNAE、NBS分类影响提案与后续付款"]),
    ("关键判断", BLUE2, ["“方案可以做”与“主体能这样签、团队能这样交”是两回事", "经营范围、设备安排与合同范围必须放在一起处理，不能只转给律师"]),
    ("本人行动", BLUE, ["现场牵头LTC团队，商法财到场集中处理公司设立与合同准备", "组织方案侧与财务、律师对齐服务内容与CNAE、NBS（最终核定由会计师承担）", "按纪要行动项，承担与华为沟通解决方案的责任", "推动范围、责任矩阵、验收、实施计划与付款安排一致"]),
    ("可见变化", GREEN, ["问题进入签约前提清单，责任人明确", "本地支持层级、总部参与方式、培训对象、数据与先期算力等交付前提明确", "各项外部条件分别跟踪（见下）"]),
], size=10)
sy = CT + 3.45
sw = (CW - 4 * 0.16) / 5
for i, lab in enumerate(["本地注册", "客户资金", "伙伴材料", "数据中心建设", "客户配合条件"]):
    x = L + i * (sw + 0.16)
    card(s, x, sy, sw, 0.7, fill=AMBER_T)
    text(s, x + 0.1, sy, sw - 1.0, 0.7, lab, size=10.5, color=NAVY, bold=True, valign="m", sa=0)
    chip(s, x + sw - 0.9, sy + 0.19, 0.78, 0.32, "跟踪中", fill=AMBER, size=9.5)
card(s, L, sy + 0.86, CW, 0.62, fill=TINT)
chip(s, L + 0.14, sy + 1.01, 0.8, 0.32, "归因", fill=NAVY, size=9.5)
text(s, L + 1.05, sy + 0.86, CW - 1.2, 0.62, "问题由RNP在采购讨论中提出；本人承担的是跨专业协调与推进责任。未关闭事项不写“已化解”。", size=10.5, valign="m", sa=0)
notes(s, "主讲2分钟。证明跨专业落地：识别“方案能做”与“主体能这样签、团队能这样交”的差异，并承担把问题推向解决的责任。")

# ───────────────────────── P13 SERPRO 成果 ─────────────────────────
s = content_slide("1-管理角色澄清", "SERPRO成果：从高层合作意愿，推进到具体客户应用验证", role="目标管理类")
stages = [
    ("合作入口", ["高层来访", "主权AI合作伙伴框架RFI阶段材料提交", "与华为联合驻场开展场景规划"]),
    ("场景筛选", ["移动医疗车影像辅助诊断", "教育快赢场景", "政务人事管理；组织产线需求交流与方案匹配"]),
    ("需求收敛", ["SERPRO导入MGI-SGP招聘、档案数字化需求", "完成需求分析，组织Demo准备与演示", "6月招聘场景获确认开展PoC"]),
    ("具体成果", ["8月：招聘应用PoC完成开发并持续优化", "过程演示反馈积极", "进入云环境部署及验收条件沟通"]),
]
gap = 0.24
w = (CW - 3 * gap) / 4
for i, (t, b) in enumerate(stages):
    x = L + i * (w + gap)
    card(s, x, CT, w, 2.5, fill=TINT if i < 3 else TINT2)
    dot(s, x + 0.16, CT + 0.16, 0.36, str(i + 1), fill=BLUE if i < 3 else GREEN)
    text(s, x + 0.6, CT + 0.14, w - 0.7, 0.4, t, size=13, color=NAVY, bold=True, valign="m", sa=0)
    text(s, x + 0.12, CT + 0.66, w - 0.24, 1.75, [(k, {"bullet": True}) for k in b], size=11, sa=4)
    if i < 3:
        chevron(s, x + w + 0.04, CT + 0.18, fill=BLUE2)
cy = CT + 2.72
titled_card(s, L, cy, cw2, 1.25, "本人贡献", ["需求、方案、客户、资源与节奏的持续牵引；开发属研发团队产出"], title_size=12)
titled_card(s, L + cw2 + 0.24, cy, cw2, 1.25, "工作价值", ["巴西合作不只停留在国家大模型建设，开始形成面向政府业务部门的应用验证基础"], title_size=12)
note(s, cy + 1.35, "状态：PoC验收、采购路径与合同尚未完成，见下页。")
notes(s, "主讲1分钟。归因清楚：开发是研发产出，我的贡献是需求、方案、客户、资源和节奏的牵引。")

# ───────────────────────── P14 SERPRO 落地链条 ─────────────────────────
s = content_slide("1-管理角色澄清", "SERPRO落地难点：需求确认、开发验证、云资源与采购路径同步推进", role="目标管理类 ｜ 备讲")
table(s, L, CT, [2.2, 5.4, 4.93], ["推进链条", "已开展工作", "仍需收口的事项"], [
    ["需求与客户确认", "Demo、调研清单、PoC方案及每周例会", "SERPRO确认与最终业务客户SGP确认不能混同；局部需求变化需重新对齐"],
    ["开发与验证", "协调产线完成场景开发、过程演示与优化", "依据最终确认条件完成测试与验收"],
    ["云资源与部署", "推动华为云资源申请、系统部署及后续资源问题协调", "资源供给与交付节奏存在外部依赖"],
    ["商业采购", "讨论云市场上架与既有MGI—SERPRO—华为框架衔接；形成香港主体或巴西代理等方案", "最终路径、上架、预算及合同尚需继续落实"],
], row_h=1.0, size=10.5)
bar(s, CT + 0.42 + 4.0 + 0.3, "没有把Demo成功等同于项目成功：同步组织客户确认、产品验证、部署资源与采购安排。", h=0.55)
notes(s, "备讲页。周报已记录SERPRO确认与SGP确认的区别、资源等待、开发优化及销售背景导入。")

# ───────────────────────── P15 团队从0到1 ─────────────────────────
s = content_slide("1-管理角色澄清", "团队从0到1：从招募人员到形成重点区域主责配置", role="成果二 ｜ 人与团队类")
cols3 = [
    ("人员配置", "2025年底启动招募与首轮带教；人员分散，临时跨区支撑", "Q1记录共6人，拉美、中东中亚、东南亚、北部非洲、南部非洲均配置主责咨询"),
    ("责任分工", "本人直接处理更多方案细节与客户问题", "骨干承担主责方案与现场工作；本人保留关键判断、复核与总体推进"),
    ("工作方法", "按单次任务交付，经验留在个人", "集中办公、实战复盘、关键材料评审、跨项目经验复用"),
]
w3 = (CW - 2 * 0.26) / 3
for i, (t, before, after) in enumerate(cols3):
    x = L + i * (w3 + 0.26)
    text(s, x, CT, w3, 0.38, t, size=13, color=NAVY, bold=True, valign="m", sa=0)
    card(s, x, CT + 0.44, w3, 0.82, fill="F4F6F9")
    chip(s, x + 0.12, CT + 0.55, 0.6, 0.28, "起点", fill=GRAYC, size=9)
    text(s, x + 0.8, CT + 0.44, w3 - 0.9, 0.82, before, size=10.5, color=MUTE, valign="m", sa=0)
    box(s, x + w3 / 2 - 0.16, CT + 1.32, 0.32, 0.22, fill=BLUE2, shape=MSO_SHAPE.DOWN_ARROW)
    card(s, x, CT + 1.6, w3, 1.0, fill=TINT)
    chip(s, x + 0.12, CT + 1.76, 0.6, 0.28, "现在", fill=BLUE, size=9)
    text(s, x + 0.8, CT + 1.6, w3 - 0.9, 1.0, after, size=10.5, color=TXT, valign="m", sa=0)
ry = CT + 2.85
text(s, L, ry, 3.0, 0.32, "区域主责与代表任务", size=11, color=NAVY, bold=True, valign="m", sa=0)
regions = [("拉美", "巴西RNP、SERPRO"), ("中东中亚", "KNG、IGNITE、MOI、哈萨克斯坦"), ("东南亚", "SDEC、泰国、斯里兰卡"), ("北部非洲", "埃及ISD"), ("南部非洲", "加纳、埃塞俄比亚、坦桑尼亚、肯尼亚、南非")]
rw = (CW - 4 * 0.16) / 5
for i, (rg, pj) in enumerate(regions):
    x = L + i * (rw + 0.16)
    card(s, x, ry + 0.38, rw, 0.98, fill=TINT)
    text(s, x + 0.1, ry + 0.42, rw - 0.2, 0.34, rg, size=11.5, color=BLUE, bold=True, valign="m", sa=0)
    text(s, x + 0.1, ry + 0.76, rw - 0.2, 0.58, pj, size=9.5, color=TXT, valign="t", sa=0)
cy = ry + 1.5
card(s, L, cy, CW, 0.62, fill=AMBER_T)
chip(s, L + 0.14, cy + 0.15, 1.0, 0.32, "实际约束", fill=AMBER, size=9.5)
text(s, L + 1.25, cy, CW - 1.4, 0.62, "人员离职、双现场常驻与轮换、小颗粒商机并发，调度压力持续；本人同步推动招聘、优先级取舍与任务交接。", size=10.5, valign="m", sa=0)
note(s, cy + 0.7, "口径：人数与区域按Q1阶段记录，不写成当前人数或整个任期累计；区域与国家对应关系按部门实际配置核对。")
notes(s, "主讲1分钟。不是接手成熟团队，而是在业务已经发生、需求持续进入的情况下建立能被部署并承担任务的咨询力量。")

# ───────────────────────── P16 骨干培养 ─────────────────────────
s = content_slide("1-管理角色澄清", "骨干培养：在重大项目实战中，把关键能力交给团队", role="人与团队类")
lw = 5.6
text(s, L, CT, lw, 0.36, "徐唯：责任递进（由本人招募并持续带教）", size=12, color=NAVY, bold=True, valign="m", sa=0)
steps = ["参与材料准备与客户交流", "承担巴西主要方案与客户问题应答", "现场主责咨询；本人转向关键判断、复核与总体推进"]
for i, st in enumerate(steps):
    y = CT + 0.46 + i * 0.98
    card(s, L, y, lw, 0.82, fill=TINT if i < 2 else TINT2)
    dot(s, L + 0.16, y + 0.23, 0.36, str(i + 1), fill=BLUE if i < 2 else NAVY)
    text(s, L + 0.66, y, lw - 0.8, 0.82, st, size=11, color=TXT, valign="m", sa=0)
    if i < 2:
        box(s, L + lw / 2 - 0.12, y + 0.84, 0.24, 0.12, fill=BLUE2, shape=MSO_SHAPE.DOWN_ARROW)
rx = L + lw + 0.3
rw = R - rx
table(s, rx, CT, [rw / 2, rw / 2], ["现在能独立承担", "仍由本人把关"], [
    ["一般方案材料与客户问题书面应答", "重大效果承诺与技术路线选择"],
    ["现场主责咨询与日常客户沟通", "工作范围、权利边界与对外正式承诺"],
    ["联合准备、面客与会后复盘", "交易模式、报价与付款设计"],
    ["主责交接与新人带教", "跨部门资源协调与升级事项"],
], row_h=0.62, size=10, first_col_bold=False)
text(s, rx, CT + 0.42 + 4 * 0.62 + 0.04, rw, 0.3, "清单以本人确认为准", size=9, color=MUTE, valign="m", sa=0)
my = CT + 3.5
mech = [
    ("授权边界", "一般方案交给骨干；重大效果、范围、权利、对外承诺继续把关，避免“完全放手”与“全部自己做”两个极端"),
    ("连续性管理", "人员阶段性离场时，本人接续关键工作并跨区调配；随后恢复分工，不以个人无限补位替代组织承接"),
    ("带教延伸到销售", "4月：对新到现场销售导入项目背景、政府与客户关系，明确现场目标；8月持续导入项目信息与伙伴资源"),
]
for i, (t, b) in enumerate(mech):
    x = L + i * (w3 + 0.26)
    card(s, x, my, w3, 1.35, fill=TINT)
    text(s, x + 0.14, my + 0.1, w3 - 0.28, 0.34, t, size=11.5, color=NAVY, bold=True, valign="m", sa=0)
    text(s, x + 0.14, my + 0.44, w3 - 0.28, 0.9, b, size=10, color=TXT, valign="t", sa=0)
text(s, L, my + 1.48, CW, 0.4, "组织承接能力已在形成；对本人的依赖尚未消除，是下一阶段的管理重点。", size=11, color=NAVY, bold=True, valign="m", sa=0)
notes(s, "主讲1.5分钟。重点讲右表：骨干现在能独立承担什么、哪些仍由我把关。这比“成长很快”具体。")

# ───────────────────────── P17 埃及 ─────────────────────────
s = content_slide("1-管理角色澄清", "埃及ISD：持续连接客户咨询、PoC、合同与交付准备", role="运营类 ｜ 备讲")
bh = 1.5
band(s, L, CT, CW, bh, "本人判断与组织", [
    "向铁三角与交付侧交底前期方案、背景与约束", "识别原HLD部分内容未完全响应RFP，提示客户可能挑战",
    "与法务合规讨论范围风险；审阅合同初版并提出修改", "提出主合同不足影响与集成商合同的背靠背衔接"], color=BLUE)
band(s, L, CT + bh + 0.14, CW, bh, "团队与专业产出", [
    "AI Lab负责人与技术团队交流，开展RAG现状调查", "现场技术方案与演示准备",
    "PoC测试与演示跟进，参与复盘并提出下一步建议", "安全许可、本地测试效果等约束下的现场协调"], color=BLUE2)
band(s, L, CT + 2 * (bh + 0.14), CW, 1.2, "已发生的阶段变化", [
    "形成基于HLD的交付安排", "合同修改意见提交主责咨询", "售前方案、客户需求、合同责任与交付执行之间未出现断层"], color=GREEN, cols=1)
bar(s, CT + 2 * (bh + 0.14) + 1.2 + 0.22, "咨询工作没有止步于客户展示，而是持续防止售前方案、客户需求、合同责任与交付执行之间出现断层。", h=0.55)
notes(s, "备讲页。追问埃及时使用。")

# ───────────────────────── P18 中东中亚 ─────────────────────────
s = content_slide("1-管理角色澄清", "中东中亚：围绕竞争、范围与交付条件开展咨询支撑", role="岗位特性类 ｜ 备讲")
cases = [
    ("IGNITE（巴基斯坦）", ["新版招标文件存在无法满足项；伙伴方案与外部竞争影响参与方式", "讨论平台、语料、训练指导等最小可参与范围"], ["组织内部研判，向一线反馈策略"]),
    ("KNG（科威特）", ["客户仍按初版SoC理解范围，尚未最终接受新HLD", "推动客户、华为与总集继续确认"], ["SoC评审；联合产线评估交付风险"]),
    ("沙特", ["MOI硬件与接口问题需其他团队协同", "SDAIA智能体及平台机会持续跟进"], ["协同组织；PoC受阻事项如实列为风险与处理工作"]),
]
for i, (t, j, o) in enumerate(cases):
    x = L + i * (w3 + 0.26)
    card(s, x, CT, w3, 2.85, fill=TINT)
    text(s, x + 0.14, CT + 0.1, w3 - 0.28, 0.4, t, size=12.5, color=NAVY, bold=True, valign="m", sa=0)
    chip(s, x + 0.14, CT + 0.58, 0.6, 0.28, "判断", fill=BLUE, size=9)
    text(s, x + 0.12, CT + 0.9, w3 - 0.24, 1.4, [(k, {"bullet": True}) for k in j], size=10.5, sa=4)
    chip(s, x + 0.14, CT + 1.95, 0.6, 0.28, "产出", fill=GREEN, size=9)
    text(s, x + 0.12, CT + 2.27, w3 - 0.24, 0.55, [(k, {"bullet": True, "bcolor": GREEN}) for k in o], size=10.5, sa=4)
oy = CT + 3.05
card(s, L, oy, CW, 1.15, fill="F4F6F9")
text(s, L + 0.14, oy + 0.08, 2.0, 0.32, "其他产出", size=11, color=NAVY, bold=True, valign="m", sa=0)
others = [("哈萨克斯坦", "MaaS与Agent平台客户汇报、华为赋能"), ("科威特", "政务、知识库及配置材料"), ("阿曼", "油气方案交流"), ("白俄罗斯", "客户接待与技术沟通")]
ow = (CW - 0.28 - 3 * 0.16) / 4
for i, (k, v) in enumerate(others):
    x = L + 0.14 + i * (ow + 0.16)
    text(s, x, oy + 0.42, ow, 0.3, k, size=10.5, color=BLUE, bold=True, valign="m", sa=0)
    text(s, x, oy + 0.7, ow, 0.4, v, size=9.5, color=TXT, valign="t", sa=0)
bar(s, oy + 1.3, "不是机械响应销售需求，而是基于竞争、能力与交付条件，给出可以继续推进的方案与边界。", h=0.55)
notes(s, "备讲页。KNG可作为全球线的追问案例：客户认知与合同范围不一致的风险管理。")

# ───────────────────────── P19 东南亚 ─────────────────────────
s = content_slide("1-管理角色澄清", "东南亚：从应用方案支撑，走向总集边界、预算与采购材料", role="目标管理类")
lw = 8.75
text(s, L, CT, lw, 0.36, "SDEC（东马）", size=12.5, color=NAVY, bold=True, valign="m", sa=0)
bands = [
    ("本人判断与组织", BLUE, ["华为提出讯飞作为总集并集成其他伙伴产品：识别伙伴交付、商务安排与利润达标风险，组织内部策略讨论", "识别员工助手方案匹配卡点，协调一网通牵头方案整合", "安排咨询赴现场推进确认"], 1.72),
    ("团队与专业产出", BLUE2, ["农业数字人、员工助手等第二轮客户方案输出", "HLD编写"], 1.0),
    ("已发生的阶段变化", GREEN, ["8月完成范围与SoC细节对齐，拆为5个子项目", "相应SoC发送客户，等待商务与客户流程"], 1.0),
]
y = CT + 0.44
for lab, col, items, h in bands:
    band(s, L, y, lw, h, lab, items, cols=1, color=col, label_w=1.55)
    y += h + 0.12
rx = L + lw + 0.28
rw = R - rx
titled_card(s, rx, CT, rw, 1.95, "泰国", ["预算约束下，调整平台与知识转移方案及报价，使客户期望与可投入资源匹配"], title_size=12)
titled_card(s, rx, CT + 2.1, rw, 1.65, "斯里兰卡", ["二期援助方案初版"], title_size=12)
bar(s, y + 0.1, "从“做应用方案”延伸到总集范围、内部产品匹配、利润风险与采购文件，判断项目的可执行性。", h=0.55)
notes(s, "主讲1.5分钟。先讲为什么需要管理判断（总集责任、伙伴集成、利润），再讲产出（五个子项目SoC）。")

# ───────────────────────── P20 非洲 ─────────────────────────
s = content_slide("1-管理角色澄清", "非洲多国：形成国别方案、语料与运营设计等实际产出", role="岗位特性类 ｜ 备讲")
table(s, L, CT, [1.55, 4.9, 3.25, 2.83], ["国家 / 机会", "具体工作产出", "阶段价值与边界", "本人作用"], [
    ["加纳", "语料调研策略、运营方案、HLD、配置报价更新；后续技术建议书、语料本地化及应用差异化说明", "从通用介绍进入本地资源、运营与实施方案；预算与推进待跟踪", "制定调研策略；复用巴西经验"],
    ["埃塞俄比亚", "伙伴选择RFI初版、方案与概算评审、预算范围沟通；8月与华为就基础方案达成一致", "形成客户与伙伴讨论的具体输入，不等同正式合同", "把关材料；统一伙伴口径"],
    ["坦桑尼亚", "概算输出及HLD持续更新", "阶段性建设与投入讨论材料", "评审与经验复用"],
    ["肯尼亚", "可研材料持续修改与内部评审", "支撑项目可行性讨论", "内部评审"],
    ["南非、加蓬等", "高阶训战、招采场景PoC初步沟通、智算中心项目沟通", "新机会输入，按前期阶段列示", "协调人员与伙伴口径"],
], row_h=0.86, size=10)
note(s, 6.15, "材料由主责咨询与团队完成；本人作用为策略制定、关键把关与经验复用，不写成个人主笔。", h=0.4)
notes(s, "备讲页。")

# ───────────────────────── P21 资源经营 ─────────────────────────
s = content_slide("1-管理角色澄清", "资源经营：不只增加商机覆盖，也判断哪些值得投入", role="目标管理类")
items = [
    ("识别低成熟度机会", "无当年预算、无明确需求的线索不应同等投入", "4月周报提出减少低价值消耗，优先聚焦可落地项目"),
    ("调整进入路径", "沙特STC当年模型训练机会不足", "提出优先小型应用或人才培训，不按原大项目方向持续投入"),
    ("控制方案规模", "客户期望需与可投入资源匹配", "IGNITE讨论最小化工作范围；泰国按预算调整方案与报价"),
    ("优先处理高风险接口", "总集、未确认范围、合同输入不完整最容易失控", "提前组织评审，不以“材料交出去了”作为工作结束"),
]
ch = 2.15
for i, (t, j, a) in enumerate(items):
    x = L + (i % 2) * (cw2 + 0.24)
    y = CT + (i // 2) * (ch + 0.2)
    card(s, x, y, cw2, ch, fill=TINT)
    dot(s, x + 0.16, y + 0.16, 0.34, str(i + 1))
    text(s, x + 0.6, y + 0.12, cw2 - 0.7, 0.42, t, size=12.5, color=NAVY, bold=True, valign="m", sa=0)
    chip(s, x + 0.16, y + 0.72, 0.6, 0.28, "判断", fill=BLUE2, size=9)
    text(s, x + 0.84, y + 0.66, cw2 - 1.0, 0.6, j, size=10.5, color=TXT, valign="m", sa=0)
    chip(s, x + 0.16, y + 1.38, 0.6, 0.28, "动作", fill=BLUE, size=9)
    text(s, x + 0.84, y + 1.32, cw2 - 1.0, 0.7, a, size=10.5, color=TXT, valign="m", sa=0)
bar(s, CT + 2 * ch + 0.42, "咨询人力有限时，提出更有针对性的投入策略；没有数据支撑，不宣称成本或转化率变化。", h=0.55)
notes(s, "主讲1分钟。讲总监的取舍，不重复项目列表。")

# ───────────────────────── P22 方案资产 ─────────────────────────
s = content_slide("1-管理角色澄清", "方案资产与产品反馈：把单个项目经验用于其他国家", role="运营类 ｜ 备讲")
tops = [
    ("国别方案", "巴西、埃塞俄比亚、巴基斯坦、加纳；按数据、预算、能力建设与场景持续更新"),
    ("跨国复用", "巴西、加纳经验用于IGNITE方案讨论；加纳与巴西方案基线拉通"),
    ("统一口径", "开闭源模型选择等问题，整理现场反馈形成团队标准表达材料"),
    ("产品反馈", "参与开源大模型方案研发开工、三大中心国际版研发讨论"),
]
w4 = (CW - 3 * 0.2) / 4
for i, (t, b) in enumerate(tops):
    x = L + i * (w4 + 0.2)
    card(s, x, CT, w4, 1.5, fill=TINT)
    text(s, x + 0.14, CT + 0.1, w4 - 0.28, 0.34, t, size=12, color=NAVY, bold=True, valign="m", sa=0)
    text(s, x + 0.14, CT + 0.46, w4 - 0.28, 1.0, b, size=10, color=TXT, valign="t", sa=0)
table(s, L, CT + 1.72, [2.3, 3.0, 4.2, 3.03], ["材料", "本人作用", "使用项目", "复用情况"], [
    ["国别方案", "组织框架、把关判断", "巴西、埃塞俄比亚、巴基斯坦、加纳", "跨国复用基线"],
    ["HLD", "评审与关键技术判断", "巴西RNP、加纳、埃及、SDEC", "结构复用"],
    ["SoC", "范围与交付风险评审", "科威特KNG、SDEC", "方法复用"],
    ["技术问题答复", "亲自完成关键澄清", "巴西RNP（4批30多个问题）", "口径沉淀"],
    ["演示与汇报材料", "组织与复核", "SERPRO Demo、哈萨克斯坦汇报", "场景复用"],
], row_h=0.56, size=10)
note(s, CT + 1.72 + 0.42 + 5 * 0.56 + 0.1, "主笔与采用情况以确认记录为准；未完成标准化的部分列入改进计划。")
notes(s, "备讲页。")

# ───────────────────────── P23 伙伴赋能 ─────────────────────────
s = content_slide("1-管理角色澄清", "伙伴赋能与高层交流：让公司能力进入客户与伙伴的决策讨论", role="岗位特性类 ｜ 备讲")
items = [
    ("华为伙伴赋能", ["覆盖拉美、南部非洲、中亚及相关代表处", "主权AI、平台与场景方案交流，面向伙伴销售与方案人员"]),
    ("重大活动分享", ["MWC圆桌、北非智算周", "分别标注本人分享、团队准备或本人统筹，不统一写“参加”"]),
    ("高层客户交流", ["SERPRO高层、哥伦比亚相关代表团及其他政企客户", "围绕客户需求输出方案，形成后续事项"]),
    ("集团能力拉通", ["金光集团接待中统筹大模型板块", "拉通工业智能、企业数字化、金融等产线准备材料、演示与汇报"]),
]
ch = 1.95
for i, (t, b) in enumerate(items):
    x = L + (i % 2) * (cw2 + 0.24)
    y = CT + (i // 2) * (ch + 0.2)
    titled_card(s, x, y, cw2, ch, t, b, num=i + 1, size=11)
bar(s, CT + 2 * ch + 0.42, "把分散的产品能力组织成面向客户的整体价值表达；活动次数不直接等同商机收入。", h=0.55)
notes(s, "备讲页。")

# ───────────────────────── P24 工作机制 ─────────────────────────
s = content_slide("1-管理角色澄清", "工作机制与提效：已开展实践，下一步提升稳定性与复用", role="运营类 ｜ 备讲")
ch = 2.15
# 卡1 集中办公（含数据）
x, y = L, CT
card(s, x, y, cw2, ch, fill=TINT)
dot(s, x + 0.16, y + 0.16, 0.34, "1")
text(s, x + 0.6, y + 0.12, cw2 - 0.7, 0.42, "集中办公与带教", size=12.5, color=NAVY, bold=True, valign="m", sa=0)
stat(s, x + 0.16, y + 0.66, 1.75, 0.8, "3次 · 18天", "合肥集中办公", fill=WHITE)
stat(s, x + 2.05, y + 0.66, 1.75, 0.8, "4天", "广州骨干集中办公", fill=WHITE)
text(s, x + 0.16, y + 1.55, cw2 - 0.32, 0.55, [("常态化辅导；5月北京集中办公与现场评审安排", {"bullet": True})], size=10.5, sa=0)
rest = [
    ("协同接口", ["处理销售与咨询协同、跨过咨询直接对接后端等问题", "参与LTC协同梳理与作战地图宣贯"]),
    ("关键材料复核", ["国别方案、标书、SoC、报价与合同范围评审", "发现问题后组织对应专业人员，不留到签约前"]),
    ("AI辅助Demo探索", ["5月底探索快速构建项目Demo，6月开展试点", "已有AI Connection Demo等产出；不虚构提效比例"]),
]
for i, (t, b) in enumerate(rest, start=1):
    x = L + (i % 2) * (cw2 + 0.24)
    y = CT + (i // 2) * (ch + 0.2)
    titled_card(s, x, y, cw2, ch, t, b, num=i + 1, size=11)
bar(s, CT + 2 * ch + 0.42, "咨询组织不只靠临时派人与个人经验运行，已开始建设共同的协作与工作方法；成熟度仍需提升。", h=0.55)
notes(s, "备讲页。")

# ───────────────────────── P25 角色自评汇总 ─────────────────────────
s = content_slide("1-管理角色澄清", "各项展开述职：四类角色的完成情况与可完善之处", role="承上启下")
table(s, L, CT, [1.85, 5.7, 3.25, 1.73], ["管理角色", "完成较好之处与关键举措", "可完善之处", "自评"], [
    ["岗位特性类", "关键咨询亲自上场；技术路线获认可并进入TR；多国国别方案形成（P8、P9、P20）", "标准口径材料覆盖不全，部分判断仍依赖本人", ("达成", {"bold": True, "color": GREEN})],
    ["目标管理类", "政府间文件推进、交易准备、SERPRO PoC、商机取舍（P7、P10、P13、P21）", "商业闭环未完成；前置事项可更早明确", ("阶段达成", {"bold": True, "color": AMBER})],
    ["人与团队类", "主责配置、骨干承接、销售导入（P15、P16）", "单点依赖仍在；主备机制待完善", ("形成中", {"bold": True, "color": AMBER})],
    ["运营类", "采购与合同书面化、本地承接矛盾处理、SoC与合同评审、机制建设（P11、P12、P17、P24）", "标准化与资产管理不足", ("达成", {"bold": True, "color": GREEN})],
], row_h=0.9, size=10.5, aligns=["l", "l", "l", "c"])
ly = CT + 0.42 + 4 * 0.9 + 0.3
lenses = [("看定位", "是否准确 → P3 总判断"), ("看结果", "是否达到预期 → P6 阶段成果与边界"), ("看过程动作", "是否扎实到位 → P7、P8、P12 案例四层"), ("看时间精力", "是否充分保障 → P5 双线并发")]
for i, (k, v) in enumerate(lenses):
    x = L + i * (w4 + 0.2)
    card(s, x, ly, w4, 0.8, fill=TINT2)
    text(s, x + 0.12, ly + 0.05, w4 - 0.24, 0.32, k, size=11, color=BLUE, bold=True, valign="m", sa=0)
    text(s, x + 0.12, ly + 0.37, w4 - 0.24, 0.4, v, size=9.5, color=TXT, valign="t", sa=0)
notes(s, "主讲1分钟。把案例收回四类角色，回应模板“各项展开述职”的要求，并过渡到不足分析。")

# ───────────────────────── P26 不足分析 ─────────────────────────
s = content_slide("2-不足分析与改进", "不足分析与改进：商业闭环、组织韧性与成果管理仍需加强")
card(s, L, CT, CW, 1.75, fill=TINT)
dot(s, L + 0.16, CT + 0.16, 0.34, "1")
text(s, L + 0.6, CT + 0.12, 6, 0.42, "商业闭环仍未完成", size=12.5, color=NAVY, bold=True, valign="m", sa=0)
subs = [("经营状态", GRAYC, "正式合同、回款、最终验收尚待推进"), ("本人可改进", BLUE, "前置事项更早明确、资源更早配置、专业评审前移"), ("外部依赖", AMBER, "客户审批、资金、伙伴与政府程序：需要管理，但不能单方面保证")]
sw = (CW - 0.32 - 2 * 0.2) / 3
for i, (k, c, v) in enumerate(subs):
    x = L + 0.16 + i * (sw + 0.2)
    card(s, x, CT + 0.62, sw, 1.0, fill=WHITE)
    chip(s, x + 0.12, CT + 0.72, 1.05, 0.3, k, fill=c, size=9.5)
    text(s, x + 0.12, CT + 1.05, sw - 0.24, 0.55, v, size=10.5, color=TXT, valign="t", sa=0)
rows = [
    ("关键知识仍集中", "部分客户背景、重大判断依赖本人与少数骨干", "完善主备、背景交接、问题台账与升级机制"),
    ("成果状态与投入口径待统一", "各区域成果状态、投入与客户反馈的口径不一，影响资源配置与效果评估", "建立项目级成果台账：日期、产出、客户反馈、当前状态"),
    ("复用与标准化仍不充分", "已有跨项目复用，但部分材料与方法仍依赖个人组织", "形成版本责任、适用条件、复用反馈与更新机制"),
]
rh = 1.02
for i, (t, cur, imp) in enumerate(rows):
    y = CT + 1.95 + i * (rh + 0.14)
    card(s, L, y, CW, rh, fill=TINT)
    dot(s, L + 0.16, y + 0.16, 0.34, str(i + 2))
    text(s, L + 0.6, y + 0.12, 3.3, 0.42, t, size=12, color=NAVY, bold=True, valign="m", sa=0)
    chip(s, L + 4.0, y + 0.18, 0.9, 0.3, "当前表现", fill=GRAYC, size=9.5)
    text(s, L + 4.0, y + 0.5, 3.9, 0.5, cur, size=10, color=TXT, valign="t", sa=0)
    chevron(s, L + 8.05, y + 0.35, fill=BLUE2)
    chip(s, L + 8.35, y + 0.18, 0.9, 0.3, "改进方向", fill=BLUE, size=9.5)
    text(s, L + 8.35, y + 0.5, 4.0, 0.5, imp, size=10, color=TXT, valign="t", sa=0)
notes(s, "主讲1.5分钟。不把“尚未签约”全部归为个人不足，也不回避缺口。第一项拆成经营状态、本人可改进、外部依赖三类。")

# ───────────────────────── P27 巴西规划 ─────────────────────────
s = content_slide("3-工作规划与路径", "巴西下一阶段：将战略成果转化为合同、交付与持续经营基础")
plans = [
    ("RNP签约", ["以范围、验收、责任、付款、签约主体与客户采购条件为主线，收口待决事项", "政府间文件与预算事项持续协同"]),
    ("交付衔接", ["把合同承诺转成交付计划", "明确数据、算力、专家、培训对象、客户配合与变更机制"]),
    ("SERPRO应用", ["完成最终业务客户确认、部署与PoC验收", "推进确定的采购路径、预算及合同讨论"]),
    ("本地承接", ["落实实体、专业服务、本地协调、开票与支持安排", "逐步形成项目后续履约基础"]),
]
for i, (t, b) in enumerate(plans):
    x = L + i * (w4 + 0.2)
    titled_card(s, x, CT, w4, 2.3, t, b, num=i + 1, size=11)
card(s, L, CT + 2.5, CW, 1.0, fill=TINT2)
chip(s, L + 0.14, CT + 2.65, 1.0, 0.32, "所需资源", fill=NAVY, size=9.5)
text(s, L + 1.25, CT + 2.5, CW - 1.4, 1.0, [("关键节点的法财、算法、产品及服务支持", {"bullet": True}), ("现场主备安排，以及明确的跨部门决策接口", {"bullet": True})], size=10.5, valign="m", sa=2)
note(s, CT + 3.6, "以上为拟推进计划，不混入已完成成果；不预设已获得完整巴西经营授权。")
notes(s, "主讲半分钟。")

# ───────────────────────── P28 咨询规划 ─────────────────────────
s = content_slide("3-工作规划与路径", "国际咨询下一阶段：提高骨干独立性与重点商机转化效率")
plans = [
    ("人员与分工", ["按重点项目成熟度配置主责与备份", "减少长期单点依赖"]),
    ("重点商机", ["按客户预算、决策阶段、方案差异与实施条件分级投入", "定期复核继续、调整或停止"]),
    ("方案资产", ["国别方案、核心技术问题、SoC、验收与合同风险案例统一管理", "明确责任人与版本"]),
    ("交付前置", ["总集、私有化部署、非标准功能、模型训练等高风险内容", "让专业部门更早介入"]),
]
for i, (t, b) in enumerate(plans):
    x = L + i * (w4 + 0.2)
    titled_card(s, x, CT, w4, 2.3, t, b, num=i + 1, size=11)
card(s, L, CT + 2.5, CW, 1.0, fill=TINT2)
chip(s, L + 0.14, CT + 2.65, 1.0, 0.32, "管理检验", fill=NAVY, size=9.5)
text(s, L + 1.25, CT + 2.5, CW - 1.4, 1.0, [("跟踪材料被采用、客户确认、阶段推进、骨干独立承接与跨项目复用情况", {"bullet": True}), ("不只统计会议与文件数量；具体目标值由本人与组织确定", {"bullet": True})], size=10.5, valign="m", sa=2)
notes(s, "主讲半分钟。")

# ───────────────────────── P29 自我发展 ─────────────────────────
s = content_slide("4-自我发展", "自我发展：从高强度亲自攻坚，走向更稳定地组织结果")
devs = [
    ("经营判断", "把方案、范围、成本、付款与风险放进同一判断框架", "深入参与重点项目报价、合同与交付复盘，与财务、商务共同检验假设"),
    ("授权与育人", "提高骨干独立决策与承接能力，保留关键事项把关", "明确授权边界、复核节点与主备安排，以真实项目进行带教"),
    ("跨组织协同", "从靠个人持续催办，转向清晰的责任与决策机制", "把重复出现的问题形成责任接口、书面记录与升级安排"),
]
for i, (t, g, p) in enumerate(devs):
    x = L + i * (w3 + 0.26)
    card(s, x, CT, w3, 3.1, fill=TINT)
    dot(s, x + 0.16, CT + 0.16, 0.36, str(i + 1))
    text(s, x + 0.62, CT + 0.14, w3 - 0.74, 0.4, t, size=13, color=NAVY, bold=True, valign="m", sa=0)
    chip(s, x + 0.16, CT + 0.78, 0.9, 0.3, "发展目标", fill=BLUE2, size=9.5)
    text(s, x + 0.16, CT + 1.12, w3 - 0.32, 0.6, g, size=11, color=TXT, valign="t", sa=0)
    chip(s, x + 0.16, CT + 1.8, 0.9, 0.3, "实践方式", fill=BLUE, size=9.5)
    text(s, x + 0.16, CT + 2.14, w3 - 0.32, 0.9, p, size=11, color=TXT, valign="t", sa=0)
note(s, CT + 3.25, "对应下一阶段工作的实际要求，不泛泛写“加强沟通、提升领导力”。")
notes(s, "主讲半分钟。")

# ───────────────────────── P30 课程学习自检 ─────────────────────────
s = content_slide("5-课程学习自检", "课程学习自检")
table(s, L + 1.5, CT + 0.2, [4.5, 2.0, 3.03], ["课程名称", "课程类型", "学习 / 考试结果"], [
    ["讯飞干部应知应会课程（名称按学习平台记录填写）", "必修", "以学习平台记录为准"],
    ["干部管理课程（名称按学习平台记录填写）", "必修", "以学习平台记录为准"],
    ["", "选修", ""],
    ["", "选修", ""],
], row_h=0.6, size=10.5, aligns=["l", "c", "c"], first_col_bold=False)
note(s, CT + 0.2 + 0.42 + 4 * 0.6 + 0.2, "说明：课程与考试明细以学习平台记录为准；参加干部培训不等同于相应课程全部通过。")
notes(s, "按学习平台记录补充课程与考试结果。")

# ───────────────────────── P31 建议 ─────────────────────────
s = content_slide("6-其他建议", "对公司和部门的建议：把重大项目实践转化为组织能力")
sug = [
    ("战略项目责任机制", ["明确总体负责人、客户接口、专业责任人与重大承诺决策链", "使总体责任与必要的协调权限匹配"]),
    ("出海专业支撑前置", ["本地签约主体、税务开票、用工、软件权利、伙伴合同等共性问题", "建立可复用的专业支撑与升级安排，避免每个项目临近签约时重新探索"]),
    ("咨询投入与激励", ["资源配置与商机成熟度、战略价值、交付可行性匹配", "关注长期驻外、轮换与跨时区工作的保障，提升队伍连续承接能力"]),
]
for i, (t, b) in enumerate(sug):
    x = L + i * (w3 + 0.26)
    titled_card(s, x, CT, w3, 2.6, t, b, num=i + 1, title_size=13, size=11.5)
note(s, CT + 2.75, "基于实战提出的组织建议，不表述为已经获批的新制度或权限。")
notes(s, "主讲1分钟。")

# ───────────────────────── 内容区垂直居中（标题下方与页脚之间） ─────────────────────────
def vcenter(slide, target_bottom=6.8):
    shapes = [sh for sh in slide.shapes if not sh.is_placeholder and sh.top / 914400 >= CT - 0.05 and sh.top / 914400 < 6.95]
    if not shapes:
        return
    bottom = max((sh.top + sh.height) / 914400 for sh in shapes)
    if bottom >= target_bottom - 0.35:
        return
    delta = (target_bottom - bottom) / 2
    for sh in shapes:
        sh.top = sh.top + Inches(delta)


for sl in list(prs.slides)[10:]:
    vcenter(sl)

# ───────────────────────── 删除模板说明页并排序 ─────────────────────────
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
keep_tpl = {0, 1, 9}
for i in sorted(range(10), reverse=True):
    if i not in keep_tpl:
        prs.part.drop_rel(ids[i].rId)
        sldIdLst.remove(ids[i])
ids = list(sldIdLst)
cover_id, toc_id, end_id = ids[0], ids[1], ids[2]
new_ids = ids[3:]
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in [cover_id, toc_id] + new_ids + [end_id]:
    sldIdLst.append(el)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
