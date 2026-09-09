# -*- coding: utf-8 -*-
"""
v9 替换页（第二批）：只做主汇报第 5 页与第 9 页，供手工替换进 v8。
  第 1 页 = 主汇报第 5 页“国家 AI 应用”：补 RNP 参与产品定义与功能设计、交付标准产品、交付后自主运营与演进。
  第 2 页 = 主汇报第 9 页“知识转移的四个阶段”：跟岗补真实训练机会，正式移交补面向运维团队与生产环境交接。
版式与 v8 完全一致；课时与人数未改（等徐唯的新数字）。本文件不做页内跳转链接。
用法：python3 build_v9_p5_p9.py <template.pptx> <out.pptx>
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from pptx_helpers import (BLUET, BODY, CW, L, LINE, MUTE, PANEL, R, RED, REDT, TXT, WHITE,
                          arrow, body, box, card, chip, drop_first_slide, label, new_slide, text)

TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else "模板-华为简报.pptx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "v9-替换页-第5页与第9页.pptx"
RNP_C = "1F4E79"
XF, RN, BOTH = REDT, BLUET, PANEL

prs = Presentation(TEMPLATE)
LAYOUT = prs.slide_layouts[0]


def slide(title, size=22, ref=None):
    s = new_slide(prs, LAYOUT, title, size=size)
    if ref:
        x = R
        for code in reversed([c.strip() for c in ref.split("·")]):
            x -= 1.08
            chip(s, x, 0.32, 1.0, 0.3, "附录 " + code, fill=WHITE, color=RNP_C, size=9.5, line=LINE)
    return s


def thesis(s, content, y=1.05, h=0.62):
    body(s, L, y, CW, h, content, size=14, color=TXT, valign="m", ls=1.25)


def legend(s, y, x=None):
    x = R - 4.1 if x is None else x
    for i, (f, t) in enumerate([(REDT, "讯飞主导"), (BLUET, "RNP 主导"), (PANEL, "双方")]):
        bx = x + i * 1.35
        box(s, bx, y + 0.06, 0.22, 0.16, fill=f, line=LINE)
        text(s, bx + 0.28, y, 1.0, 0.28, t, size=9.5, color=MUTE, valign="m", margin=(0, 0, 0, 0), sa=0)


def steps(s, y, h, items, gap=0.25, size=12):
    n = len(items)
    bw = (CW - (n - 1) * gap) / n
    for i, (name, desc, fill) in enumerate(items):
        bx = L + i * (bw + gap)
        card(s, bx, y, bw, h, fill=fill)
        text(s, bx + 0.2, y + 0.14, bw - 0.4, 0.4, f"{i + 1}  {name}", size=13, color=TXT, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
        body(s, bx + 0.2, y + 0.62, bw - 0.4, h - 0.7, desc, size=size, ls=1.3)
        if i < n - 1:
            arrow(s, bx + bw + 0.02, y + 0.22, fill=LINE)


def two_boxes(s, y, h, xf, rnp, size=12, rnp_title="RNP 怎么参与"):
    half = (CW - 0.3) / 2
    for i, (name, color, fill, items) in enumerate([("讯飞负责", RED, REDT, xf), (rnp_title, RNP_C, BLUET, rnp)]):
        cx = L + i * (half + 0.3)
        card(s, cx, y, half, h, fill=fill)
        label(s, cx + 0.25, y + 0.1, 3.0, name, color=color, size=13)
        body(s, cx + 0.25, y + 0.5, half - 0.5, h - 0.56, items, size=size, bullet=True, bcolor=color, sa=3)


def means(s, y, content, h=0.6):
    card(s, L, y, CW, h, fill=WHITE, line=RNP_C)
    body(s, L + 0.25, y + 0.05, CW - 0.5, h - 0.1, "**对 RNP 意味着**　" + content, size=12.5, color=RNP_C, valign="m")


# ============================== 第 5 页 国家 AI 应用（替换） ==============================
s = slide("国家 AI 应用：讯飞交付软件，RNP 负责内容、用户与运营", size=21, ref="A4")
thesis(s, "国民 APP（手机、电脑多终端）和数据运营平台，面向 RNP 批准的用户；提供问答、检索、文档阅读、语音对话、写作")
legend(s, 1.78)
steps(s, 2.1, 1.6, [
    ("明确需求", "功能、用户范围、内容规则双方共同确定；RNP 参与产品定义与功能设计", BOTH),
    ("本地化", "在标准产品上做葡语界面与本地化配置，采用 RNP 的品牌和内容", XF),
    ("系统集成", "RNP 提供本地系统接口、账号凭证与搜索等第三方服务，讯飞完成集成", BOTH),
    ("上线运营", "Stage 1 初步交付，Stage 2 扩展运行；之后由 RNP 运营并自主演进", RN),
], size=11.5)
two_boxes(s, 3.95, 1.45,
          ["交付标准产品与本地化配置；接通接口、配置平台", "APP 演示已在第 3 页平台演示中安排"],
          ["参与产品定义与功能设计，提供内容、品牌与本地系统接口",
           "审批用户；上线后由产品、运营和审核团队承接，可在交付版本上自主演进"])
means(s, 5.65, "容量按 1,000 路并发问答配置并实测；产品团队在需求阶段即需到位，运营团队在 Stage 1 交付前到位")

# ============================== 第 9 页 知识转移四阶段（替换） ==============================
s = slide("知识转移的四个阶段：578 课时、约 74 个工作日", size=21, ref="A8")
thesis(s, "一支核心团队，移交后能独立完成约定的平台运维、训练提交、模型部署与发布；30 人建立基础，10 人深入到独立操作")
legend(s, 1.78)
phases = [
    ("奠基", "30", "约 5 个工作日 · 30 人",
     "大模型基础、项目架构、语料到服务全流程、安全与运维基础", XF),
    ("跟岗", "360", "约 45 个工作日 · 10 人",
     "讯飞操作，RNP 观察协助；覆盖国家模型训练、12 个月内的模型升级与 3 项 DLM 训练等真实任务", XF),
    ("指导实操", "20", "约 3 个工作日 · 10 人",
     "RNP 操作：数据处理、训练提交、模型部署、监控与故障处理；双方签实操记录", RN),
    ("正式移交", "168", "约 21 个工作日 · 10 人",
     "面向 RNP 运维团队：交接生产环境、运维流程、资产许可与运行手册；验证独立操作，双方签确认", RN),
]
gap = 0.25
pw = (CW - 3 * gap) / 4
py, ph = 2.1, 2.35
for i, (name, hours, meta, tasks, fill) in enumerate(phases):
    px = L + i * (pw + gap)
    card(s, px, py, pw, ph, fill=fill)
    text(s, px + 0.2, py + 0.12, pw - 0.4, 0.36, f"{i + 1}  {name}", size=13, color=TXT, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
    text(s, px + 0.2, py + 0.5, pw - 0.4, 0.46, [f"**{hours}** 课时"], size=12, color=TXT, valign="m", margin=(0, 0, 0, 0), sa=0)
    r0 = s.shapes[-1].text_frame.paragraphs[0].runs[0]
    r0.font.size = Pt(22)
    r0.font.color.rgb = RGBColor.from_string(RED)
    text(s, px + 0.2, py + 0.98, pw - 0.4, 0.3, meta, size=11, color=MUTE, valign="m", margin=(0, 0, 0, 0), sa=0)
    body(s, px + 0.2, py + 1.32, pw - 0.4, ph - 1.37, tasks, size=10.5, ls=1.25)
two_boxes(s, 4.7, 1.15,
          ["由参与本项目交付的讯飞工程师带教，检查关键步骤", "移交材料与独立操作验证"],
          ["指定稳定、具备基础的学员，阶段学员可重合", "安排时间，参与真实交付任务"], size=11.5)
means(s, 6.05, "74 个工作日是培训活动天数，分散在实施期内，不是连续工期；跟岗随真实训练任务安排")

drop_first_slide(prs)
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
