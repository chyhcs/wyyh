# -*- coding: utf-8 -*-
"""
讯飞 × RNP 技术交流 PPT（中文版，10 页），基于华为模板生成。
内容已按 Mota 电话、9 月 4 日内部会议和 12 页框架填好；
蓝色文字为需要按 SOW、报价、技术方案核实或替换的内容。

用法：python3 build_cn_deck.py <template.pptx> <out.pptx>
文本标记：{{...}} 内的文字渲染为蓝色（待替换）。
"""
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else "template.pptx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "out.pptx"

# ───────────── 设计令牌（沿用模板：红色标题、灰色面板、微软雅黑） ─────────────
FONT = "微软雅黑"
RED = "C00000"        # 模板标题红
RED2 = "E9002F"       # 主题 accent1
TXT = "1D1D1A"
BODY = "333333"
MUTE = "666666"
LINE = "D9D9D9"
PANEL = "F2F2F2"
PANEL2 = "F7F7F7"
WHITE = "FFFFFF"
BLUE = "0070C0"       # 待替换标记色
GREEN = "61B230"      # 已包含
GREEN_T = "EAF5E1"
ORANGE = "ED6D00"     # 可选
ORANGE_T = "FDEBDA"
GRAY = "8C8C8C"       # 不在范围 / 待确认
GRAY_T = "EFEFEF"
REDT = "FBE5E9"       # 讯飞栏淡红
BLUET = "E8F0FA"      # RNP 栏淡蓝

W, H = 13.333, 7.5
L, R = 0.5, 12.83
CW = R - L

prs = Presentation(TEMPLATE)
LAYOUT = prs.slide_layouts[0]  # "简报"：只有一个标题占位符


def rgb(h):
    return RGBColor.from_string(h)


# ───────────── 文本 ─────────────
def _runs(s):
    """'{{蓝色}}' 标记 → (text, is_blue, is_bold)；'**粗**' → 粗体。"""
    out = []
    for part in re.split(r"(\{\{.*?\}\})", s):
        if not part:
            continue
        blue = part.startswith("{{") and part.endswith("}}")
        if blue:
            part = part[2:-2]
        for sub in re.split(r"(\*\*.*?\*\*)", part):
            if not sub:
                continue
            bold = sub.startswith("**") and sub.endswith("**")
            if bold:
                sub = sub[2:-2]
            out.append((sub, blue, bold))
    return out


def _set_font(r, size, color, bold):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        for e in rPr.findall(qn(tag)):
            rPr.remove(e)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)


def _bullet(p, color, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.15))))
    pPr.set("indent", str(-int(Inches(0.15))))
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def fill_tf(tf, content, size=10.5, color=BODY, bold=False, align="l", valign="t",
            margin=(0.05, 0.03, 0.05, 0.03), ls=1.1, sa=2, bullet=False, bcolor=None):
    """content: str（'\n' 分段）或 list[str]。"""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[valign]
    txBody = tf._txBody
    for p in list(txBody.findall(qn("a:p"))):
        txBody.remove(p)
    paras = content if isinstance(content, list) else content.split("\n")
    for line in paras:
        p = tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = ls
        p.space_after = Pt(sa)
        for t, blue, b in _runs(line):
            r = p.add_run()
            r.text = t
            _set_font(r, size, BLUE if blue else color, bold or b)
        if bullet:
            _bullet(p, bcolor or color)
    return tf


def text(slide, x, y, w, h, content, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, content, **kw)
    return tb


# ───────────── 形状 ─────────────
def box(slide, x, y, w, h, fill=None, line=None, radius=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def card(slide, x, y, w, h, fill=PANEL, line=None, radius=0.04):
    return box(slide, x, y, w, h, fill=fill, line=line, radius=radius, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def chip(slide, x, y, w, h, label, fill=RED, color=WHITE, size=9.5, line=None):
    shp = box(slide, x, y, w, h, fill=fill, line=line, radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    fill_tf(shp.text_frame, label, size=size, color=color, bold=True, align="c", valign="m", margin=(0.04, 0, 0.04, 0), sa=0)
    return shp


def dot(slide, x, y, d, label, fill=RED2, color=WHITE, size=10):
    shp = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    fill_tf(shp.text_frame, label, size=size, color=color, bold=True, align="c", valign="m", margin=(0, 0, 0, 0), sa=0)
    return shp


def arrow(slide, x, y, w=0.2, h=0.2, fill=LINE):
    return box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.RIGHT_ARROW)


def label(slide, x, y, w, content, color=TXT, size=12, h=0.3):
    return text(slide, x, y, w, h, content, size=size, color=color, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)


def body(slide, x, y, w, h, content, size=10.5, color=BODY, **kw):
    return text(slide, x, y, w, h, content, size=size, color=color, margin=(0, 0, 0, 0), **kw)


def note(slide, y, content, h=0.36, size=9.5):
    return text(slide, L, y, CW, h, content, size=size, color=MUTE, valign="m", margin=(0, 0, 0, 0), sa=0)


# ───────────── 表格 ─────────────
def cell_borders(cell, color=LINE, w=0.5):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(Pt(w))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(i, ln)


def table(slide, x, y, colw, header, rows, row_h=0.6, header_h=0.4, size=9.5, header_size=10,
          header_fill=RED, header_color=WHITE, row_fills=None, first_col=True, first_col_fill=None, valign="t", pad=0.06):
    nrows, ncols = len(rows) + 1, len(colw)
    heights = row_h if isinstance(row_h, list) else [row_h] * len(rows)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(colw)), Inches(header_h + sum(heights)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(colw):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, nrows):
        tbl.rows[r].height = Inches(heights[r - 1])
    for c, h in enumerate(header):
        cell = tbl.cell(0, c)
        cell_borders(cell, color=header_fill if header_color == WHITE else LINE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(header_fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_tf(cell.text_frame, h, size=header_size, color=header_color, bold=True, valign="m", margin=(pad, 0.02, pad, 0.02), sa=0)
    for r, row in enumerate(rows, start=1):
        rf = row_fills[r - 1] if row_fills else WHITE
        for c, content in enumerate(row):
            cell = tbl.cell(r, c)
            cell_borders(cell)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(first_col_fill if (c == 0 and first_col_fill) else rf)
            cell.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[valign]
            bold = first_col and c == 0
            fill_tf(cell.text_frame, content, size=size, color=TXT if bold else BODY, bold=bold,
                    valign=valign, margin=(pad, 0.04, pad, 0.04), sa=1, ls=1.08)
    return tbl


# ───────────── 页面骨架 ─────────────
PAGE = [0]


def new_slide(title, size=22):
    s = prs.slides.add_slide(LAYOUT)
    PAGE[0] += 1
    for ph in s.placeholders:
        # 模板占位符：红色、粗体、微软雅黑；这里只改字号和文字
        fill_tf(ph.text_frame, title, size=size, color=RED, bold=True, valign="b", margin=(0.05, 0, 0.05, 0), sa=0)
    return s


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def two_party(slide, x, y, w, h, left, right, size=9.5):
    """讯飞 / RNP 两栏分工。"""
    half = (w - 0.12) / 2
    card(slide, x, y, half, h, fill=REDT)
    card(slide, x + half + 0.12, y, half, h, fill=BLUET)
    label(slide, x + 0.1, y + 0.06, half - 0.2, "讯飞", color=RED, size=10.5, h=0.26)
    label(slide, x + half + 0.22, y + 0.06, half - 0.2, "RNP", color="1F4E79", size=10.5, h=0.26)
    body(slide, x + 0.1, y + 0.36, half - 0.2, h - 0.42, left, size=size)
    body(slide, x + half + 0.22, y + 0.36, half - 0.2, h - 0.42, right, size=size)


# =====================================================================
# 第 1 页 封面
# =====================================================================
s = new_slide("", size=22)
text(s, 0.7, 1.35, 8.4, 0.4, "RNP × 科大讯飞 · 技术交流", size=13, color=MUTE, bold=True, margin=(0, 0, 0, 0))
text(s, 0.7, 1.8, 8.4, 1.6, "葡语大模型、应用与领域模型", size=34, color=RED, bold=True, valign="t", margin=(0, 0, 0, 0), ls=1.05)
text(s, 0.7, 2.75, 8.4, 0.6, "讯飞交付范围、共同实施与知识转移方案", size=20, color=TXT, bold=True, margin=(0, 0, 0, 0))
chip(s, 0.7, 3.65, 1.7, 0.32, "讨论稿 v0.1", fill=RED)
text(s, 0.7, 4.15, 8.6, 1.5, [
    "会议：RNP 技术交流 · 里约热内卢 · {{2026 年 9 月 9 日（周三）}}",
    "汇报：科大讯飞 陈浩、{{方案架构师姓名}}",
    "参会：RNP  Yara 及技术团队；华为  路易斯、郑丽雪、{{董阳光}}",
    "本材料内容以 SOW、报价与技术方案核对结果为准",
], size=12, color=BODY, margin=(0, 0, 0, 0), ls=1.3, sa=4)
# 图例
card(s, 0.7, 5.9, 8.4, 0.62, fill=BLUET)
text(s, 0.85, 5.9, 8.1, 0.62, "图例：{{蓝色文字}} = 待按 SOW、报价、技术方案核实或替换的内容（日期、人名、数字、技术路线、范围口径）",
     size=10.5, color=BODY, valign="m", margin=(0, 0, 0, 0), sa=0)
# 右侧示意：三项交付 + 知识转移
bx, by, bw, bh, g = 9.65, 1.75, 3.0, 0.88, 0.2
for i, t in enumerate(["葡语基础模型 LLM", "应用 APP", "领域模型 DLM"]):
    card(s, bx, by + i * (bh + g), bw, bh, fill=PANEL)
    dot(s, bx + 0.2, by + i * (bh + g) + (bh - 0.36) / 2, 0.36, str(i + 1))
    text(s, bx + 0.7, by + i * (bh + g), bw - 0.8, bh, t, size=14, color=TXT, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
card(s, bx, by + 3 * (bh + g), bw, 0.7, fill=RED)
text(s, bx, by + 3 * (bh + g), bw, 0.7, "项目内知识转移 · 共同实施", size=12.5, color=WHITE, bold=True, align="c", valign="m", margin=(0, 0, 0, 0), sa=0)
notes(s, "本页作用：确定本次交流范围——讯飞的模型、应用、领域模型工作，双方共同实施，以及知识转移。让 RNP 一开始就知道后面既谈交付成果，也谈本方如何组织人员参与。\n"
         "需填：日期、汇报人、版本状态。\n"
         "依据：Mota 电话 01:03–02:29（解释 LLM/APP/DLM 怎么做、双方责任、巴方要安排人并写自己的计划）。")

# =====================================================================
# 第 2 页 交付成果总览
# =====================================================================
s = new_slide("讯飞交付成果总览：四项成果各自的用途与巴方参与")
cards = [
    ("葡语基础模型 LLM",
     "葡语大模型，含{{模型权重、推理服务}}、评测报告与操作文档",
     "政务场景的葡语问答、文本生成与理解；作为 APP 与领域模型的共同底座",
     "提供葡语语料并落实授权；参与评测与验收；指定联合团队",
     "已包含", GREEN, GREEN_T),
    ("应用 APP",
     "面向政务人员的葡语智能应用，含账号权限、模型调用与业务系统对接",
     "{{政策咨询问答、公文辅助起草、材料摘要}}",
     "提供业务系统接口与测试数据；确认业务流程并组织验收",
     "已包含", GREEN, GREEN_T),
    ("领域模型 DLM",
     "{{1–2 个}}领域模型及配套应用与文档",
     "解决{{教育、医疗或公共安全}}领域的具体业务问题",
     "选定领域；提供领域数据与专家；参与评价与验收",
     "已包含 · {{领域数量待确认}}", GREEN, GREEN_T),
    ("项目内知识转移",
     "讲解、示范、共同操作与指导实践；配套资料与实践记录",
     "巴方团队掌握数据处理、模型训练调优、应用开发与运行",
     "按角色投入人员承担实际任务，{{各阶段名额按 SOW}}",
     "已包含在报价 · {{需核对}}", GREEN, GREEN_T),
]
gap = 0.25
cw = (CW - 3 * gap) / 4
cy, ch = 1.05, 5.15
for i, (name, deliv, use, rnp, st, sc, st_t) in enumerate(cards):
    cx = L + i * (cw + gap)
    card(s, cx, cy, cw, ch, fill=PANEL2, line=LINE)
    dot(s, cx + 0.18, cy + 0.18, 0.36, str(i + 1))
    text(s, cx + 0.62, cy + 0.12, cw - 0.75, 0.48, name, size=13, color=TXT, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
    yy = cy + 0.78
    for lab, val in (("交付物", deliv), ("用途", use), ("巴方参与", rnp)):
        label(s, cx + 0.18, yy, cw - 0.36, lab, color=RED, size=10.5, h=0.26)
        body(s, cx + 0.18, yy + 0.28, cw - 0.36, 0.9, val, size=10)
        yy += 1.22
    card(s, cx + 0.18, cy + ch - 0.6, cw - 0.36, 0.4, fill=st_t)
    text(s, cx + 0.26, cy + ch - 0.6, cw - 0.52, 0.4, "状态：" + st, size=9.5, color=sc, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
note(s, 6.35, "材料中的案例、功能示例、中间版本与本项目交付物分别标注；每项成果在第 3–5、8 页展开。")
notes(s, "本页作用：先回答最基本的问题——讯飞这部分最终形成什么、各有什么用、巴方要做什么。让后面每页有明确位置。\n"
         "要求：巴方参与要能看出实际工作，不能只写参加培训。\n"
         "材料：讯飞技术方案、SOW 与交付清单；展示材料要标明是案例、中间版本还是正式交付。\n"
         "衔接：第 3–5 页展开 LLM/APP/DLM，第 8 页展开知识转移。")

# =====================================================================
# 第 3 页 葡语基础模型
# =====================================================================
s = new_slide("葡语基础模型：交付什么、如何共同建设")
label(s, L, 1.0, 4, "建设路径", color=TXT, size=12)
steps = [
    ("数据准备", "收集与清洗葡语语料；巴方提供政务语料并落实授权"),
    ("适配 / 训练", "{{以讯飞星火底座做葡语增量预训练与指令微调}}"),
    ("评测", "葡语通用基准 + 政务任务评测集，双方共同评审"),
    ("交付", "模型包、推理服务、评测报告、操作文档与验收"),
]
gap = 0.3
bw = (CW - 3 * gap) / 4
fy, fh = 1.35, 1.1
for i, (name, hint) in enumerate(steps):
    bx = L + i * (bw + gap)
    card(s, bx, fy, bw, fh, fill=REDT if i == 0 else PANEL)
    dot(s, bx + 0.12, fy + 0.12, 0.3, str(i + 1), size=9.5)
    text(s, bx + 0.5, fy + 0.08, bw - 0.6, 0.38, name, size=11.5, color=TXT, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
    body(s, bx + 0.12, fy + 0.52, bw - 0.24, fh - 0.58, hint, size=9.5)
    if i < 3:
        arrow(s, bx + bw + 0.05, fy + fh / 2 - 0.1)

cy, ch = 2.75, 2.75
cw = (CW - 2 * gap) / 3
cols = [
    ("目标能力与衡量方式", [
        "支持{{葡语政务问答、公文起草、信息抽取与摘要}}",
        "评测：葡语通用基准 + {{N 个}}政务任务评测集，达标线{{待定}}",
        "交付前提供样例对比与评测报告，供 RNP 评审",
    ]),
    ("数据与条件", [
        "来源：公开葡语语料 + RNP 及政府部门提供的政务语料，{{规模待定}}",
        "授权与脱敏由巴方负责，数据处理由讯飞负责",
        "依赖：华为算力与训练环境按计划到位，{{算力规模待定}}",
    ]),
]
for i, (h, items) in enumerate(cols):
    cx = L + i * (cw + gap)
    card(s, cx, cy, cw, ch, fill=PANEL2, line=LINE)
    label(s, cx + 0.18, cy + 0.14, cw - 0.36, h, color=RED, size=11.5)
    body(s, cx + 0.18, cy + 0.52, cw - 0.36, ch - 0.65, items, size=10, bullet=True, bcolor=RED, sa=5)
cx = L + 2 * (cw + gap)
card(s, cx, cy, cw, ch, fill=PANEL2, line=LINE)
label(s, cx + 0.18, cy + 0.14, cw - 0.36, "双方分工", color=RED, size=11.5)
two_party(s, cx + 0.18, cy + 0.52, cw - 0.36, ch - 0.7,
          ["数据处理与清洗", "模型训练与调优", "评测与报告", "打包、文档与交付"],
          ["需求确认", "语料供给与授权", "参与评测评审", "组织验收"], size=9.5)

card(s, L, 5.72, CW, 0.78, fill=PANEL)
label(s, L + 0.18, 5.78, 3, "交付与验收", color=RED, size=11)
body(s, L + 0.18, 6.08, CW - 0.36, 0.4,
     "交付形式：{{模型权重 + 推理服务镜像}}  ·  验收：评测达标 + 试运行{{N 周}}  ·  配套：训练数据说明、评测报告、操作指南", size=10)
notes(s, "本页作用：把第 2 页的“葡语 LLM”展开，让 RNP 明白模型怎样形成、讯飞负责什么、巴方要提供和参与什么。同时服务成果理解和实施安排。\n"
         "口径：不预设必须从零训练；不展开技术选型依据，听众需要的是阶段、分工和结果。模型能力与数据处理都要有明确承担方。\n"
         "材料：模型方案、数据要求、评测材料、交付形式、双方任务说明。\n"
         "依据：Mota 电话（怎么做、责任）；9 月 4 日会议（讲结果，客户不查技术栈）。")

# =====================================================================
# 第 4 页 应用 APP
# =====================================================================
s = new_slide("应用 APP：谁来用、解决什么问题")
lx, ly, lw, lh = L, 1.05, 5.3, 4.35
card(s, lx, ly, lw, lh, fill=PANEL, line=LINE)
text(s, lx + 0.3, ly + lh / 2 - 0.55, lw - 0.6, 0.45, "界面截图 · 演示视频 · 真实操作流程", size=14, color=MUTE, bold=True, align="c", valign="m", margin=(0, 0, 0, 0), sa=0)
text(s, lx + 0.3, ly + lh / 2 - 0.05, lw - 0.6, 0.45, "{{此处插入本项目 APP 或参考案例的界面与端到端流程}}", size=10.5, color=BODY, align="c", valign="m", margin=(0, 0, 0, 0), sa=0)
card(s, lx + 0.3, ly + lh - 0.62, lw - 0.6, 0.4, fill=WHITE)
text(s, lx + 0.4, ly + lh - 0.62, lw - 0.8, 0.4, "标注材料性质：现有案例 / 功能示例 / 本项目交付", size=9.5, color=MUTE, valign="m", margin=(0, 0, 0, 0), sa=0)

rx = lx + lw + 0.4
rw = R - rx
rows = [
    ("用户与场景", "{{政府部门业务人员}}；场景：{{政策咨询、公文办理、材料审阅}}"),
    ("功能与结果", "{{葡语对话问答、文档摘要与审阅、公文模板生成、知识库检索}}；结果：可直接使用的文稿与答复"),
    ("规模与性能", "{{支持 N 用户、并发 N、平均响应 N 秒}}，以双方确认值为准"),
    ("对应的政务用例", "{{参考中国政务案例：××部门智能办公助手}}（标注为案例，非本项目交付）"),
]
for i, (h, t) in enumerate(rows):
    ry = 1.05 + i * 1.1
    dot(s, rx, ry + 0.03, 0.34, str(i + 1))
    label(s, rx + 0.48, ry, rw - 0.5, h, color=TXT, size=12)
    body(s, rx + 0.48, ry + 0.36, rw - 0.5, 0.7, t, size=10.5)

card(s, L, 5.65, CW, 0.85, fill=REDT)
label(s, L + 0.18, 5.72, 4, "巴方为应用接入提供的条件", color=RED, size=11)
body(s, L + 0.18, 6.02, CW - 0.36, 0.45,
     "业务系统接口  ·  测试数据  ·  测试与业务确认  ·  接入范围待确认：葡语本地化、账号权限、模型调用方式、业务系统对接边界", size=10)
notes(s, "本页作用：让 RNP 从使用者角度理解交付——谁用、完成什么工作、得到什么结果。功能要放进真实使用过程，听众不需要技术栈。\n"
         "必须展示：一段完整操作过程的界面或视频，并标明性质；功能与结果；规模与性能（按确认值）；对应的政务用例。\n"
         "保留一行接入条件，便于 RNP 技术团队安排系统、数据和测试（Mota：讲清 APP 怎么做、谁做什么）。\n"
         "依据：9 月 4 日会议 34:06–34:46（能接多少人、哪些功能、有视频或图片、解决哪些政务应用；可拿中国案例照着讲）。")

# =====================================================================
# 第 5 页 领域模型 DLM
# =====================================================================
s = new_slide("领域模型 DLM：围绕真实业务问题共同开发")
card(s, L, 1.02, CW, 0.58, fill=PANEL)
label(s, L + 0.18, 1.08, 1.6, "候选领域", color=TXT, size=12)
body(s, L + 1.7, 1.15, 7.8, 0.4, "{{教育、医疗、公共安全}}——按业务问题、数据可用性与专家投入选定，{{领域数量与技术路线按实际方案确定}}", size=10.5)
chip(s, R - 2.6, 1.16, 2.45, 0.3, "状态：{{候选 / 已确认}}", fill=WHITE, color=TXT, line=LINE)

steps = ["业务问题", "数据与领域专家", "开发与适配", "评价与验收", "更新与维护"]
xf = [("与业务部门共同定义任务与成功标准", "提出业务问题，指定业务负责人"),
      ("提出数据规格与标注要求", "提供领域数据与授权，安排领域专家"),
      ("基于葡语底座做领域适配与应用开发", "技术人员参与开发，共同操作"),
      ("出具评测报告，修复问题", "专家评价，业务验收"),
      ("移交更新方法与工具，支持期内支持", "接手后自行更新与维护")]
gap = 0.25
bw = (CW - 4 * gap) / 5
fy, fh = 1.85, 0.6
for i, name in enumerate(steps):
    bx = L + i * (bw + gap)
    card(s, bx, fy, bw, fh, fill=REDT if i == 0 else PANEL)
    dot(s, bx + 0.1, fy + 0.14, 0.32, str(i + 1), size=9.5)
    text(s, bx + 0.5, fy, bw - 0.58, fh, name, size=11, color=TXT, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
    if i < 4:
        arrow(s, bx + bw + 0.03, fy + fh / 2 - 0.1)
    ry1 = fy + fh + 0.18
    card(s, bx, ry1, bw, 1.15, fill=REDT)
    label(s, bx + 0.12, ry1 + 0.07, bw - 0.24, "讯飞", color=RED, size=10, h=0.26)
    body(s, bx + 0.12, ry1 + 0.36, bw - 0.24, 0.75, xf[i][0], size=9.5)
    ry2 = ry1 + 1.15 + 0.12
    card(s, bx, ry2, bw, 1.15, fill=BLUET)
    label(s, bx + 0.12, ry2 + 0.07, bw - 0.24, "RNP", color="1F4E79", size=10, h=0.26)
    body(s, bx + 0.12, ry2 + 0.36, bw - 0.24, 0.75, xf[i][1], size=9.5)

oy = 5.42
label(s, L, oy, 4, "成果分别列出，不混成笼统的“解决方案”", color=TXT, size=11)
ow = (CW - 2 * 0.3) / 3
for i, (t, d) in enumerate([("模型成果", "领域模型及评测报告"), ("应用成果", "领域应用与接口"), ("资料成果", "数据说明、操作与更新指南")]):
    ox = L + i * (ow + 0.3)
    card(s, ox, oy + 0.34, ow, 0.5, fill=PANEL)
    text(s, ox + 0.15, oy + 0.34, ow - 0.3, 0.5, "**" + t + "：**" + d, size=10, color=BODY, valign="m", margin=(0, 0, 0, 0), sa=0)
note(s, 6.4, "领域是否选对、数据是否可用、专家能否参与评价，直接决定工作能否推进；这一页是双方共同实施关系的落点。")
notes(s, "本页作用：让 RNP 明白领域成果要围绕具体业务问题形成，巴方需要组织数据、专家和技术人员参与。\n"
         "要回答：拟开展的领域或选择机制；业务问题与预期成果；数据、专家等条件；讯飞的开发或适配工作；巴方的工作；成果怎样评价、验收、更新。\n"
         "口径：候选案例标明候选状态；领域数量和技术路线由实际方案确定；模型、应用、资料分别讲清。\n"
         "依据：Mota 电话（DLM 怎么做、责任）；框架第 5 页。")

# =====================================================================
# 第 6 页 共同实施计划
# =====================================================================
s = new_slide("共同实施计划：双方何时做什么")
note(s, 1.0, "一条时间线、三方任务。RNP 据此编制本方项目计划；华为平台依赖单列，避免计入讯飞责任。", size=10.5)
phases = ["{{阶段 1 · 启动与数据准备}}\n{{M1–M2}}", "{{阶段 2 · 底座训练与评测}}\n{{M3–M5}}",
          "{{阶段 3 · 应用与领域模型开发}}\n{{M5–M8}}", "{{阶段 4 · 部署与验收}}\n{{M9–M10}}", "{{阶段 5 · 交接与支持}}\n{{M11–M12}}"]
lanes = [
    ("阶段目标", ["确认需求，组建联合团队，语料到位", "完成葡语底座训练并通过评测", "完成 APP 与领域模型开发和联调", "生产部署、试运行、验收", "资产交接、能力验证，进入支持期"], WHITE),
    ("讯飞", ["需求分析、数据规格、环境准备", "数据处理、训练、评测", "应用开发、领域适配、联调", "部署、优化、验收支持", "交接资料，支持期响应"], REDT),
    ("RNP", ["指定团队，提供语料与授权", "参与数据处理与评测", "提供接口与专家，参与开发", "组织测试与验收", "接手运行，自行更新"], BLUET),
    ("华为平台依赖", ["算力与训练环境到位", "训练算力稳定供给", "部署环境与网络", "生产环境与运维衔接", "运维支持衔接"], ORANGE_T),
    ("阶段成果", ["需求确认书、数据清单", "模型包、评测报告", "应用与领域模型版本", "验收报告", "交接清单、能力评估报告"], WHITE),
    ("启动条件", ["合同生效，团队到位", "语料与算力到位", "底座通过评测", "联调通过", "验收完成"], WHITE),
]
rows = [[name] + cells for name, cells, _ in lanes]
fills = [f for _, _, f in lanes]
first_w = 1.75
rest_w = (CW - first_w) / 5
table(s, L, 1.45, [first_w] + [rest_w] * 5, [""] + phases, rows,
      row_h=[0.62, 0.68, 0.68, 0.62, 0.58, 0.55], header_h=0.6, size=9.5, header_size=9.5,
      header_fill="E6E6E6", header_color=TXT, row_fills=fills, first_col_fill=None)
note(s, 6.05, "阶段名称、数量与时间按实际方案确定；数据准备、模型与应用建设、评测、部署、知识转移之间的先后关系需在计划中体现。")
notes(s, "本页作用：回答“这些工作怎样安排在一起，双方什么时候做什么”。直接回应 Mota 电话中 RNP 要据此编制本方计划的要求；只给讯飞内部排期不够。\n"
         "三方：Mota 原话是 ours / yours / theirs。华为平台依赖（算力、环境、接口、昇腾栈）单列一行，避免 RNP 把它算到讯飞头上。\n"
         "材料：实施计划、阶段成果、双方责任、资源依赖、评审安排。\n"
         "衔接：第 7 页说明巴方要安排什么人、准备什么资源来执行这份计划。")

# =====================================================================
# 第 7 页 巴方人员与资源
# =====================================================================
s = new_slide("巴方需要准备的人员与资源")
note(s, 1.0, "把“巴方配合”拆成能够安排的事情：谁参与、承担什么、什么时候投入、需要什么基础。", size=10.5)
roles = [
    ["项目负责人", "统筹巴方资源，对接讯飞与华为", "项目管理经验，英语沟通", "全程，{{≥30% 投入}}", "{{1}}"],
    ["业务 / 领域专家", "定义业务问题，评价领域成果", "领域业务经验", "阶段 1、3、4，{{按需}}", "{{每领域 2–3}}"],
    ["数据工程师", "语料收集、清洗、标注", "Python，数据处理基础", "阶段 1–2，{{全职}}", "{{2–3}}"],
    ["模型工程师", "参与训练、评测与调优", "Python，机器学习基础", "阶段 2–3，{{全职}}", "{{2–4}}"],
    ["应用开发工程师", "参与 APP 与领域应用开发、系统对接", "软件开发经验", "阶段 3–4，{{全职}}", "{{2–3}}"],
    ["运维工程师", "部署、运行与监控", "Linux 与容器基础", "阶段 4–5，{{≥50% 投入}}", "{{1–2}}"],
]
lw = 8.05
table(s, L, 1.45, [1.55, 2.35, 1.75, 1.6, 0.8], ["角色", "项目中的任务", "应具备的基础", "参与阶段与投入", "人数"], roles,
      row_h=0.5, header_h=0.4, size=9.5, header_size=10)
note(s, 5.05, "区分承担项目任务的人员与仅参加课程的人员；角色按实际任务取舍，不预设固定组织结构。", size=9.5)
rx = L + lw + 0.35
rw = R - rx
card(s, rx, 1.45, rw, 4.25, fill=PANEL2, line=LINE)
label(s, rx + 0.2, 1.58, rw - 0.4, "资源准备", color=RED, size=12)
res = [("数据与授权", "政务语料与领域数据；授权与脱敏责任"),
       ("环境与访问", "账号、网络、测试与实验环境"),
       ("接口与测试条件", "业务系统接口、测试数据"),
       ("责任人 · 到位时间 · 未就绪影响", "{{各项待填}}，未到位时对应阶段顺延")]
for i, (h, t) in enumerate(res):
    ry = 2.0 + i * 0.9
    dot(s, rx + 0.2, ry + 0.02, 0.3, str(i + 1), size=9.5)
    label(s, rx + 0.6, ry, rw - 0.8, h, color=TXT, size=10.5, h=0.3)
    body(s, rx + 0.6, ry + 0.32, rw - 0.8, 0.5, t, size=9.5)
card(s, L, 5.85, CW, 0.5, fill=REDT)
text(s, L + 0.18, 5.85, CW - 0.36, 0.5, "本页是 RNP 编制本方参与计划的输入：角色、人数、时间；人员口径与华为{{2900 人}}能力计划对齐，由董阳光合稿时核对总人数。",
     size=10.5, color=RED, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)
notes(s, "本页作用：客户看完应能回去组织本方人员和资源，而不是只拿到课程报名要求。这是对 Mota 电话最直接的回应（they need to allocate people… how many people… write down this program）。\n"
         "要回答：按实际工作确定的角色、任务、基础（语言、编程、学历）、参与阶段与人数；数据与授权；环境与访问；接口与测试条件；提供方、到位时间、未就绪影响。\n"
         "口径：说明双方共同处理的工作，不把数据质量或实施条件笼统交给 RNP；格式与华为第一层人才表格兼容，便于合稿核对总人数。\n"
         "依据：Mota 电话 01:03–02:29；9 月 4 日会议（入门门槛、语言）。")

# =====================================================================
# 第 8 页 项目内知识转移
# =====================================================================
s = new_slide("项目内知识转移：在真实任务中边做边学")
chip(s, L, 1.0, 3.3, 0.32, "已包含在当前报价 · 不另收费", fill=GREEN)
text(s, L + 3.45, 1.0, 6, 0.32, "{{口径需与 SOW 及报价核对后定稿}}", size=10, color=BODY, valign="m", margin=(0, 0, 0, 0), sa=0)
kt = [
    ["数据工程师", "{{N}}", "语料清洗、标注、数据集构建", "讲解 · 示范 · 共同操作", "独立完成数据集构建", "提交一个可用数据集"],
    ["模型工程师", "{{N}}", "参与训练、评测、微调", "共同操作 · 指导实践", "独立完成一次微调与评测", "微调实验与评测报告"],
    ["应用开发工程师", "{{N}}", "参与应用开发与系统对接", "示范 · 指导实践", "独立开发并集成新功能", "交付一个功能模块"],
    ["运维工程师", "{{N}}", "部署、监控、故障处理", "讲解 · 共同操作", "独立完成部署与运行", "完成一次生产部署演练"],
]
table(s, L, 1.48, [1.7, 0.7, 2.55, 2.35, 2.45, 2.58], ["参与对象", "名额", "承担的项目任务", "讯飞支持方式", "能力结果", "验证方式"], kt,
      row_h=0.5, header_h=0.4, size=9.5, header_size=10)
by_, bh = 4.05, 1.65
bw = (CW - 2 * 0.3) / 3
boxes = [
    ("项目结束后巴方团队能够", ["独立完成数据处理与模型微调", "部署并运行应用", "评测并更新领域模型"], GREEN_T, "3B7A1A"),
    ("本次转移不覆盖", ["讯飞 SOW 以外的工作包", "{{国家 AI 软件栈自研（TR 中非讯飞承担的 FP 部分）}}，措辞待定"], GRAY_T, TXT),
    ("语言、资料与记录", ["{{英语授课，配葡语助教或同传}}", "资料随交付移交", "实践记录与技术成果由巴方留存"], PANEL2, TXT),
]
for i, (h, items, f, c) in enumerate(boxes):
    bx = L + i * (bw + 0.3)
    card(s, bx, by_, bw, bh, fill=f, line=LINE if f == PANEL2 else None)
    label(s, bx + 0.18, by_ + 0.12, bw - 0.36, h, color=c, size=11.5)
    body(s, bx + 0.18, by_ + 0.48, bw - 0.36, bh - 0.55, items, size=10, bullet=True, bcolor=c, sa=3)
note(s, 5.85, "与讯飞交付相关的开发、适配、测试与优化支持落实到具体任务；名额、时间、语言与技术指导投入以 SOW 为准。", h=0.5)
notes(s, "本页作用：回答“巴方人员参与项目以后，讯飞怎样帮助他们掌握工作，最后用什么证明确实掌握了”。知识转移与实际交付结合：任务、指导、实践、能力结果；课程名称和课时只是其中一部分。\n"
         "标题口径：这一层已在项目报价中（9 月 4 日会议：已报过价、包含在项目里、培训上免费）。发给客户前必须回到 SOW 与报价核对，不能扩展成“所有能力建设都已包含”，也不等于客户已接受。\n"
         "边界：正面写清学完能做什么、不能做什么。本页不能代替 TR 中整个国家软件栈的研发支持，也不默认讯飞承担全部 FP1–FP9。\n"
         "格式：表格与华为第一层课程表（对象、基础、名额、内容、效果）保持一致，合并稿才像一个 program。\n"
         "依据：9 月 4 日会议 23:13–24:16、26:07、44:41；Mota 电话（to be with you and to learn）。")

# =====================================================================
# 第 9 页 可选核心人才研修
# =====================================================================
s = new_slide("可选：核心人才深入研修")
chip(s, L, 1.0, 2.4, 0.32, "可选 · 单独报价", fill=ORANGE)
text(s, L + 2.55, 0.98, CW - 2.55, 0.36, "在项目内知识转移之外，为少数有基础的人员增加的深入研修；不替代任何已约定内容。", size=11, color=BODY, valign="m", margin=(0, 0, 0, 0), sa=0)
cells = [
    ("对象与选拔", "{{5–10 人}}，具备编程与模型基础，从项目参与人员中选拔；{{需通过入学测评}}"),
    ("技术方向", "模型训练与调优、模型蒸馏、评测体系建设；{{可按 RNP 需求增加数据工程方向}}"),
    ("真实课题", "围绕{{巴西自有葡语模型或某一领域模型}}的实际课题开展，产出可用于项目"),
    ("导师与形式", "讯飞研发导师带教；{{合肥或里约现场 + 远程}}；配套算力与实验环境"),
    ("成果与验证", "独立完成一次模型训练或微调，形成技术报告{{或论文}}；导师评审并出具能力评估"),
    ("规模 · 周期 · 地点 · 价格", "{{N 人 · N 周 · 地点 · 价格与实施条件}}"),
]
gap = 0.3
cw = (CW - 2 * gap) / 3
ch = 1.95
for i, (h, t) in enumerate(cells):
    col, row = i % 3, i // 3
    cx = L + col * (cw + gap)
    cy = 1.5 + row * (ch + 0.28)
    card(s, cx, cy, cw, ch, fill=PANEL2, line=LINE)
    dot(s, cx + 0.2, cy + 0.2, 0.34, str(i + 1), fill=ORANGE if i == 5 else RED2)
    label(s, cx + 0.66, cy + 0.2, cw - 0.85, h, color=TXT, size=12, h=0.34)
    body(s, cx + 0.2, cy + 0.7, cw - 0.4, ch - 0.85, t, size=10.5)
note(s, 6.05, "仅对既定范围以外的内容与投入收费；已约定的知识转移或研发支持不因技术难度改列为收费增项。", h=0.4)
notes(s, "本页作用：在项目内知识转移之外，回答“如果希望少数有基础的人进一步深入，需要增加什么内容和投入”。对应内部会议讨论的可选高端研修（少数人、单独定价、私人订制，将来作为项目看得见的成果）。\n"
         "要回答：适合谁、入选基础；深入方向；围绕什么课题；导师、方式、资源；预期能力和成果、怎样检查；人数、周期、地点、价格、实施条件。\n"
         "口径：相对第 8 页增加了什么要写清，客户才能判断是否购买；已约定内容不能挪到这里收费。\n"
         "依据：9 月 4 日会议 09:18–09:50、24:59–26:43、44:59。")

# =====================================================================
# 第 10 页 范围一览与确认事项
# =====================================================================
s = new_slide("范围一览与本次需确认事项")
gap = 0.3
cw = (CW - 2 * gap) / 3
cy, ch = 1.0, 1.85
cols = [
    ("已包含在当前方案", ["葡语大模型、APP、领域模型交付", "项目内知识转移", "交接包：模型、应用、文档、许可条件，支持期{{N 个月}}"], GREEN_T, "3B7A1A"),
    ("可选 · 单独报价", ["核心人才深入研修", "{{其他可选服务}}"], ORANGE_T, "B35400"),
    ("不在范围 · 第三方 · 客户提供", ["讯飞 SOW 以外的工作包", "第三方许可与服务", "巴方提供的数据、环境与接口"], GRAY_T, TXT),
]
for i, (h, items, f, c) in enumerate(cols):
    cx = L + i * (cw + gap)
    card(s, cx, cy, cw, ch, fill=f)
    label(s, cx + 0.18, cy + 0.12, cw - 0.36, h, color=c, size=11.5)
    body(s, cx + 0.18, cy + 0.48, cw - 0.36, ch - 0.55, items, size=10, bullet=True, bcolor=c, sa=3)
note(s, 2.9, "各项对应的技术方案、SOW、报价章节与版本：{{待补}}；按核实结果标明已包含、单独报价、可选或待确认。", size=9.5)
items = [
    ["1. 模型、应用与领域工作的目标与范围", "{{待对齐}}", "领域数量、评测达标线、APP 场景", "讯飞 + RNP", "技术方案、评测集", "{{会后 1 周}}"],
    ["2. 巴方参与人员、资源与准备条件", "{{待对齐}}", "各角色人数、到位时间", "RNP", "人员名单、资源清单", "{{会后 2 周}}"],
    ["3. 共同实施安排", "{{待对齐}}", "阶段划分与时间", "讯飞", "实施计划", "{{会后 1 周}}"],
    ["4. 知识转移与能力目标", "{{待对齐}}", "名额、语言、能力验证方式", "讯飞 + RNP", "SOW 名额、能力目标", "{{会后 1 周}}"],
    ["5. 可选研修是否继续细化", "{{待对齐}}", "是否需要、人数与预算", "RNP 决定", "研修方案与报价", "{{待定}}"],
    ["6. 成果交接、许可与后续支持", "{{待对齐}}", "交付形式、许可条件、支持期", "讯飞", "交接清单、许可条款", "{{会后 2 周}}"],
    ["7. 工作范围与费用归属", "{{待对齐}}", "已包含 / 可选 / 待确认", "讯飞 + 华为", "SOW 与报价核对表", "{{会后 1 周}}"],
]
table(s, L, 3.3, [3.35, 0.95, 2.75, 1.25, 2.35, 1.68], ["确认事项", "状态", "结论 / 待决问题", "承接方", "需补充材料", "完成时间"], items,
      row_h=0.38, header_h=0.36, size=9.5, header_size=10, valign="m")
notes(s, "本页作用：让交流形成实际推进——哪些已对齐、哪些未定、下一步谁补什么、什么时候反馈。也再次回应 RNP 编制本方计划的需要。\n"
         "范围块：前面先讲内容和价值（第 2–9 页），这里再谈范围与费用，RNP 才能知道每项费用对应什么工作，识别重复或遗漏。按核实结果标明已包含 / 单独报价 / 可选 / 待确认。“知识转移已在报价中”必须回到文件核对，不能扩展成“都已包含”或“客户已接受”。\n"
         "交接（并入本页）：模型、应用、配置、脚本、接口文档、操作权限、许可与迁移条件、巴方接手后的任务、交接完成判断、后续支持与联系机制。不默认全部源码开放或无限期支持。\n"
         "确认表：会前准备清单，会中记录决定，会后形成双方行动。已确认、建议方案、待讨论要区分；材料中出现某项内容不等于双方已承诺。\n"
         "依据：9 月 4 日会议 32:22–33:53（这次会相当于锁定 SOW；材料提前发；争取技术部分当场基本谈定）。")

# ───────────── 删除模板自带的示例页，只保留新建的 10 页 ─────────────
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
prs.part.drop_rel(ids[0].rId)
sldIdLst.remove(ids[0])

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
