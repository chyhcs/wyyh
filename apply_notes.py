# -*- coding: utf-8 -*-
"""
把《讲稿-会议版.md》按“## 第 N 页”写入 PPT 各页备注栏。
原有备注（材料来源）保留在讲稿之后。

用法：python3 apply_notes.py [deck.pptx] [讲稿.md]
在 build_cn_meeting.py 生成 PPT 之后运行。
"""
import re
import sys

from pptx import Presentation

DECK = sys.argv[1] if len(sys.argv) > 1 else "RNP-讯飞交流-会议版.pptx"
MD = sys.argv[2] if len(sys.argv) > 2 else "讲稿-会议版.md"

md = open(MD, encoding="utf-8").read()
scripts = {}
for m in re.finditer(r"^## 第 (\d+) 页[^\n]*\n(.*?)(?=^## |\Z)", md, re.S | re.M):
    scripts[int(m.group(1))] = m.group(2).strip()
opening = re.search(r"^## 开场[^\n]*\n(.*?)(?=^## )", md, re.S | re.M)
closing = re.search(r"^## 收尾[^\n]*\n(.*?)(?=^---|\Z)", md, re.S | re.M)

prs = Presentation(DECK)
for i, slide in enumerate(prs.slides, start=1):
    if i not in scripts:
        continue
    body = scripts[i]
    if i == 1 and opening:
        body = "【开场】\n" + opening.group(1).strip() + "\n\n【本页】\n" + body
    if i == len(prs.slides) and closing:
        body = body + "\n\n【收尾】\n" + closing.group(1).strip()
    old = slide.notes_slide.notes_text_frame.text.strip()
    slide.notes_slide.notes_text_frame.text = body + ("\n\n【材料来源备注】\n" + old if old else "")
prs.save(DECK)
print("notes written for", len(scripts), "slides ->", DECK)
