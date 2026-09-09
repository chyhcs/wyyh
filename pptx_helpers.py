# -*- coding: utf-8 -*-
"""
python-pptx 绘图辅助（华为简报模板配色）。
文本标记：{{...}} 渲染为蓝色（待填 / 待核实）；**...** 渲染为粗体。
"""
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "微软雅黑"
RED = "C00000"
RED2 = "E9002F"
TXT = "1D1D1A"
BODY = "333333"
MUTE = "666666"
LINE = "D9D9D9"
PANEL = "F2F2F2"
PANEL2 = "F7F7F7"
WHITE = "FFFFFF"
BLUE = "0070C0"
GREEN = "61B230"
GREEN_T = "EAF5E1"
ORANGE = "ED6D00"
ORANGE_T = "FDEBDA"
GRAY = "8C8C8C"
GRAY_T = "EFEFEF"
REDT = "FBE5E9"
BLUET = "E8F0FA"

W, H = 13.333, 7.5
L, R = 0.5, 12.83
CW = R - L


def rgb(h):
    return RGBColor.from_string(h)


def _runs(s):
    out = []
    for part in re.split(r"(\{\{.*?\}\})", s):
        if not part:
            continue
        blue = part.startswith("{{") and part.endswith("}}")
        if blue:
            part = part[2:-2]
        for sub in re.split(r"(\*\*.*?\*\*)", part):
            if not sub:
                continue
            bold = sub.startswith("**") and sub.endswith("**")
            if bold:
                sub = sub[2:-2]
            out.append((sub, blue, bold))
    return out


def _set_font(r, size, color, bold):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        for e in rPr.findall(qn(tag)):
            rPr.remove(e)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", FONT)


def _bullet(p, color, char="•"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.18))))
    pPr.set("indent", str(-int(Inches(0.18))))
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def _num_bullet(p, color):
    """自动编号项目符号（1. 2. 3.），用于要点列表。"""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.22))))
    pPr.set("indent", str(-int(Inches(0.22))))
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")


def fill_tf(tf, content, size=12, color=BODY, bold=False, align="l", valign="t",
            margin=(0.05, 0.03, 0.05, 0.03), ls=1.15, sa=4, bullet=False, bcolor=None, num=False):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[valign]
    txBody = tf._txBody
    for p in list(txBody.findall(qn("a:p"))):
        txBody.remove(p)
    paras = content if isinstance(content, list) else content.split("\n")
    for line in paras:
        p = tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = ls
        p.space_after = Pt(sa)
        for t, blue, b in _runs(line):
            r = p.add_run()
            r.text = t
            _set_font(r, size, BLUE if blue else color, bold or b)
        if bullet:
            (_num_bullet if num else _bullet)(p, bcolor or color)
    return tf


def text(slide, x, y, w, h, content, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, content, **kw)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, radius=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def card(slide, x, y, w, h, fill=PANEL, line=None, radius=0.04):
    return box(slide, x, y, w, h, fill=fill, line=line, radius=radius, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def chip(slide, x, y, w, h, label, fill=RED, color=WHITE, size=10, line=None):
    shp = box(slide, x, y, w, h, fill=fill, line=line, radius=0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    fill_tf(shp.text_frame, label, size=size, color=color, bold=True, align="c", valign="m", margin=(0.04, 0, 0.04, 0), sa=0)
    return shp


def dot(slide, x, y, d, label, fill=RED2, color=WHITE, size=11):
    shp = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    fill_tf(shp.text_frame, label, size=size, color=color, bold=True, align="c", valign="m", margin=(0, 0, 0, 0), sa=0)
    return shp


def arrow(slide, x, y, w=0.22, h=0.22, fill=LINE):
    return box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.RIGHT_ARROW)


def label(slide, x, y, w, content, color=TXT, size=13, h=0.34):
    return text(slide, x, y, w, h, content, size=size, color=color, bold=True, valign="m", margin=(0, 0, 0, 0), sa=0)


def body(slide, x, y, w, h, content, size=12, color=BODY, **kw):
    return text(slide, x, y, w, h, content, size=size, color=color, margin=(0, 0, 0, 0), **kw)


def note(slide, y, content, h=0.4, size=10.5):
    return text(slide, L, y, CW, h, content, size=size, color=MUTE, valign="m", margin=(0, 0, 0, 0), sa=0)


def cell_borders(cell, color=LINE, w=0.5, vlines=True):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        if tag in ("a:lnL", "a:lnR") and not vlines:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
            tcPr.insert(i, ln)
            continue
        ln.set("w", str(int(Pt(w))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(i, ln)


def table(slide, x, y, colw, header, rows, row_h=0.6, header_h=0.42, size=11.5, header_size=11.5,
          header_fill=RED, header_color=WHITE, row_fills=None, first_col=True, valign="m", pad=0.08, vlines=True):
    nrows, ncols = len(rows) + 1, len(colw)
    heights = row_h if isinstance(row_h, list) else [row_h] * len(rows)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(colw)), Inches(header_h + sum(heights)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(colw):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, nrows):
        tbl.rows[r].height = Inches(heights[r - 1])
    for c, h in enumerate(header):
        cell = tbl.cell(0, c)
        cell_borders(cell, color=header_fill if header_color == WHITE else LINE, vlines=vlines)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(header_fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_tf(cell.text_frame, h, size=header_size, color=header_color, bold=True, valign="m", margin=(pad, 0.02, pad, 0.02), sa=0)
    for r, row in enumerate(rows, start=1):
        rf = row_fills[r - 1] if row_fills else WHITE
        for c, content in enumerate(row):
            cell = tbl.cell(r, c)
            cell_borders(cell, vlines=vlines)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(rf)
            cell.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[valign]
            bold = first_col and c == 0
            fill_tf(cell.text_frame, content, size=size, color=TXT if bold else BODY, bold=bold,
                    valign=valign, margin=(pad, 0.04, pad, 0.04), sa=1, ls=1.1)
    return tbl


def new_slide(prs, layout, title, size=22):
    s = prs.slides.add_slide(layout)
    for ph in s.placeholders:
        fill_tf(ph.text_frame, title, size=size, color=RED, bold=True, valign="b", margin=(0.05, 0, 0.05, 0), sa=0)
    return s


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def drop_first_slide(prs):
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    prs.part.drop_rel(ids[0].rId)
    sldIdLst.remove(ids[0])
